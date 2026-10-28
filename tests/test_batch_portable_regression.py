"""Regressions for large uploads, expandable worksheets, and shared deadlines."""
import csv
import gc
import time
import threading
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QInputDialog

import Scraper
from app.constants import STANDARD_HEADERS
from app.core.data_manager import DataManager
from app.core.checkpoint_service import CheckpointService, CheckpointSnapshot
from app.excel.exporter import export_records
from app.excel.importer import import_records
from app.ui_qt.controller import ApplicationController
from app.ui_qt.pages.leads_page import LeadsPage
from app.ui_qt.pages.scraper_page import ScraperPage
from app.ui_qt.main_window import MainWindow


@pytest.fixture(scope='module')
def qt():
    return QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def collect_widgets():
    yield
    gc.collect()


def controller():
    return ApplicationController()


def wait_for(qt, predicate, seconds=10):
    until = time.monotonic() + seconds
    while not predicate() and time.monotonic() < until:
        qt.processEvents(); time.sleep(0.01)
    assert predicate()


@pytest.mark.parametrize('count', [61, 1000, 10000])
def test_paste_keeps_every_row(qt, count):
    c = controller(); page = LeadsPage(c)
    column = c.data_manager.headers.index('First Name')
    page.table.sortByColumn(-1, Qt.SortOrder.AscendingOrder)
    page.table.setCurrentIndex(page.proxy.index(0, column))
    QApplication.clipboard().setText('\n'.join(f'Person {i:05}\tFamily {i:05}' for i in range(count)))
    page.table.paste_clipboard()
    assert len(c.data_manager.records) == count
    assert [row['First Name'] for row in c.data_manager.records] == [f'Person {i:05}' for i in range(count)]
    assert c.data_manager.records[-1]['Last Name'] == f'Family {count-1:05}'
    assert page.model.rowCount() == count + 20
    assert '10,000' in page.count_label.text() if count == 10000 else True
    page.close()


def test_sorted_filtered_paste_uses_original_visible_targets(qt):
    c = controller(); c.data_manager.extend([{'First Name': str(i), 'Company Name': 'Keep' if i % 2 else 'Hidden'} for i in range(12)], 'test')
    page = LeadsPage(c); col = STANDARD_HEADERS.index('First Name')
    page.proxy.set_search('Keep'); page.table.sortByColumn(col, Qt.SortOrder.DescendingOrder)
    targets = [page.proxy.mapToSource(page.proxy.index(i, col)).row() for i in range(page.proxy.rowCount())]
    page.table.setCurrentIndex(page.proxy.index(0, col))
    QApplication.clipboard().setText('\n'.join(f'Changed {i}' for i in range(100)))
    page.table.paste_clipboard()
    assert [c.data_manager.records[row]['First Name'] for row in targets] == [f'Changed {i}' for i in range(6)]
    assert all(c.data_manager.records[i]['First Name'] == str(i) for i in range(0, 12, 2))
    assert len(c.data_manager.records) == 106
    assert c.data_manager.records[-1]['First Name'] == 'Changed 99'
    page.close()


def test_add_columns_button_and_roundtrip(qt, tmp_path):
    c = controller(); c.data_manager.append_record({'First Name': 'Ada'})
    page = LeadsPage(c)
    with patch.object(QInputDialog, 'getMultiLineText', return_value=('Owner\nNext action', True)):
        page.add_columns_button.click()
    assert page.model.columnCount() == 44
    assert c.data_manager.headers[:42] == list(STANDARD_HEADERS)
    page.model.setData(page.model.index(0, 42), 'Aditya')
    page.model.setData(page.model.index(0, 43), 'Call tomorrow')
    output = tmp_path / 'custom.xlsx'
    export_records(c.data_manager.to_dataframe(), output, custom_columns=c.data_manager.custom_columns)
    book = load_workbook(output)
    assert list(book.active.values)[0][-2:] == ('Owner', 'Next action')
    book.close()
    rows, _ = import_records(output, custom_columns=c.data_manager.custom_columns)
    assert rows[0]['Owner'] == 'Aditya' and rows[0]['Next action'] == 'Call tomorrow'
    page.proxy.set_search('Call tomorrow'); assert page.proxy.rowCount() == 1
    page.close()


def test_column_names_validate_atomically():
    manager = DataManager()
    for names in (['Owner', 'owner'], ['First Name'], [''], ['=Formula']):
        with pytest.raises(ValueError): manager.add_columns(names)
        assert not manager.custom_columns


@pytest.mark.parametrize('with_rows', [True, False])
def test_custom_schema_survives_checkpoint(tmp_path, monkeypatch, with_rows):
    import app.core.checkpoint_service as checkpoints
    monkeypatch.setattr(checkpoints, 'CHECKPOINT_DIR', tmp_path / 'checkpoints')
    monkeypatch.setattr(checkpoints, 'DB_PATH', tmp_path / 'history.sqlite3')
    c = ApplicationController(); c.data_manager.add_columns(['Owner'])
    if with_rows: c.data_manager.append_record({'Owner': 'Aditya'})
    c.save_checkpoint(reason='regression')
    snap = c.checkpoints.latest(); assert snap.payload['custom_columns'] == ['Owner']
    restored = ApplicationController(); restored.recover(snap)
    assert restored.data_manager.headers[-1] == 'Owner'
    if with_rows: assert restored.data_manager.records[0]['Owner'] == 'Aditya'


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_link_file_loads_1200_rows_and_other_sheets(tmp_path, extension):
    path = tmp_path / ('links.' + extension)
    links = [f'https://www.linkedin.com/in/person-{i}' for i in range(1200)]
    if extension == 'csv':
        with path.open('w', newline='') as handle:
            csv.writer(handle).writerows([['Profile']] + [[x] for x in links])
    else:
        book = Workbook(); sheet = book.active
        sheet.append(['Profile'])
        for link in links[:800]: sheet.append([link])
        sheet = book.create_sheet('More links')
        for index, link in enumerate(links[800:], 1):
            sheet.cell(index, 1, 'Open profile').hyperlink = link
        book.save(path); book.close()
    updates = []
    items = Scraper.SearchWorker._load_link_inputs(None, path, lambda d, t: updates.append((d, t)))
    assert len(items) == 1200
    assert [item.canonical for item in items] == links
    assert updates and updates[-1][0] >= 1200


def test_scraper_file_loading_and_default_full_batch(qt, tmp_path):
    path = tmp_path / 'links.csv'
    path.write_text('\n'.join(f'https://www.linkedin.com/in/person-{i}' for i in range(1000)))
    c = controller(); page = ScraperPage(c)
    page.load_input(str(path))
    assert not page.start_button.isEnabled()
    wait_for(qt, lambda: not page._file_loading)
    assert '1,000' in page.file_info.text()
    assert page.file_progress.value() == 100 and not page.max_records.isEnabled()
    fake = Mock(); fake.is_alive.return_value = True; page.worker = fake
    page.output_path.setText(str(tmp_path / 'results.xlsx'))
    page.start_processing()
    wait_for(qt, lambda: fake.send.called)
    command, payload = fake.send.call_args.args
    assert command == 'process_links' and payload['max_records'] == 0
    assert len(payload['link_inputs']) == 1000
    page.events.put(('input_total', (1000,)))
    page.events.put(('progress', (500, 250, 200, 50, 1000)))
    page._drain_events()
    assert page.progress.value() == 50
    page._processing_done(); page.workbook.wb.close(); page.timer.stop(); page.close()


def test_shared_deadline_bounds_scraper_queries_and_allows_retry(monkeypatch):
    import lookup_budget
    from business_email import BusinessEmailEnricher
    clock = [0.0]
    monkeypatch.setattr(lookup_budget.time, 'monotonic', lambda: clock[0])
    calls = []
    def search(query):
        calls.append(query)
        clock[0] += lookup_budget.timeout(6)
        return []
    web = SimpleNamespace(search=search, fetch=lambda url: None, documents={}, source_index={})
    service = BusinessEmailEnricher(web, lambda a, b: a == b)
    row = {'First Name': 'Ada', 'Last Name': 'Lovelace', 'Company Name': 'Example',
           'Company Website': 'https://example.com', 'Business Email ID': '-', 'Email Status': '-'}
    result = service.enrich(row)
    assert clock[0] <= 15 and len(calls) == 3
    assert result['Email Status'].startswith('Timed out')
    assert result['Business Email ID'] == '-' and not service.cache and not service.pattern_cache
    service.enrich(row)
    assert len(calls) == 6 and clock[0] <= 30
    assert not lookup_budget.active()


def test_search_provider_retries_share_budget(monkeypatch):
    import httpx
    from scraper_engine import lookup_budget
    from app.services.search_provider import BraveSearchProvider
    from app.core.email_pattern_discovery import PublicEmailPatternDiscoverer
    clock = [0.0]; timeouts = []
    monkeypatch.setattr(lookup_budget.time, 'monotonic', lambda: clock[0])
    monkeypatch.setattr(lookup_budget.time, 'sleep', lambda seconds: clock.__setitem__(0, clock[0] + seconds))
    class Client:
        def __init__(self, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def request(self, *args, **kwargs):
            timeouts.append(kwargs['timeout'])
            clock[0] += kwargs['timeout']
            raise httpx.ReadTimeout('Simulated timeout')
        def close(self): pass
    monkeypatch.setattr(httpx, 'Client', Client)
    from request_limits import RequestLimits
    monkeypatch.setattr('app.services.search_provider.REQUEST_LIMITS', RequestLimits())
    cache = Mock(); cache.get.return_value = None
    provider = BraveSearchProvider(api_key='test-only', retries=100)
    result = PublicEmailPatternDiscoverer(provider, pattern_cache=cache, discovery_cache=cache).discover('example.com')
    assert result.status == 'Timed Out'
    assert 0 < clock[0] <= 15.001
    assert all(t <= 3.75 for t in timeouts)
    assert timeouts
    cache.upsert.assert_not_called()


def test_timeout_stays_retryable_and_never_creates_a_candidate():
    from app.core.email_finder import EmailFinder
    from app.core.email_pattern_discovery import PatternDiscoveryResult
    cache = Mock(); cache.get.return_value = None
    row = {'First Name': 'Ada', 'Last Name': 'Lovelace', 'Company Website': 'example.com',
           'Business Email ID': '-', 'Email Status': 'Timed out (retry available)',
           'Prospect LinkedIn Link': 'https://www.linkedin.com/in/ada-example'}
    result = EmailFinder(cache).find_for_record(0, row, discovery_result=PatternDiscoveryResult(
        domain='example.com', status='Timed Out', reason='15-second budget reached'))
    assert result.status == 'Timed Out' and not result.applied
    assert result.candidate == '-'
    worker = Scraper.SearchWorker.__new__(Scraper.SearchWorker)
    worker.cache_path = 'unused'; worker.stop_event = threading.Event(); worker.enrichment_attempts = {}
    worker._remember_enrichment(row)
    assert not worker.enrichment_attempts


def test_progress_visible_until_worker_finishes(qt):
    released = threading.Event()
    with patch.object(MainWindow, '_offer_recovery'), patch.object(MainWindow, 'refresh_qwen_status'):
        window = MainWindow(); window.show()
        window.run_worker(lambda: released.wait(4), started='Reading test file…', finished='File ready')
        qt.processEvents()
        assert window.operation_progress.isVisible() and window.operation_progress.maximum() == 0
        assert not window.leads_page.table.isEnabled()
        released.set(); wait_for(qt, lambda: not window._jobs)
        assert not window.operation_progress.isVisible() and window.leads_page.table.isEnabled()
        assert window.status_message.text() == 'File ready'
        window.close(); qt.processEvents()


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_lead_import_keeps_single_column_1500_rows(tmp_path, extension):
    path = tmp_path / ('names.' + extension)
    names = [f'Person {i}' for i in range(1500)]
    if extension == 'csv':
        with path.open('w', newline='') as handle:
            csv.writer(handle).writerows([['First Name']] + [[name] for name in names])
    else:
        book = Workbook(); book.active.append(['First Name'])
        for name in names: book.active.append([name])
        book.save(path); book.close()
    rows, _ = import_records(path)
    assert [row['First Name'] for row in rows] == names
