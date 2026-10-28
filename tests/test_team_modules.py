import json
import pytest
from tests.test_team_service import team,call,commit,PASSWORD
from app.team.access import ServiceError
from app.team.modules import MODULES,OPERATION_MODULES
from app.team.service import Service


def test_every_rpc_has_an_explicit_module_gate():
    assert {name[3:] for name in dir(Service) if name.startswith('op_')} == set(OPERATION_MODULES)


def test_department_defaults_user_overrides_and_live_revocation(team):
    assert not call(team,'quality','modules')['scraper']['allowed']
    assert call(team,'a','modules')['scraper']['allowed']
    call(team,'admin','module_policy_save',subject_type='department',subject_id='operations',overrides={'scraper':False,'reports':False})
    with pytest.raises(ServiceError,match='Scraper'):call(team,'a','module_check',module='scraper')
    with pytest.raises(ServiceError,match='Reports'):call(team,'a','report',campaign_id='C1')
    call(team,'admin','module_policy_save',subject_type='user',subject_id=team[2]['a']['id'],overrides={'scraper':True})
    assert call(team,'a','modules')['scraper']=={'label':'Scraper','allowed':True,'source':'User override'}
    call(team,'a','module_check',module='scraper')
    call(team,'admin','module_policy_save',subject_type='user',subject_id=team[2]['a']['id'],overrides={'scraper':None})
    with pytest.raises(ServiceError):call(team,'a','module_check',module='scraper')
    call(team,'admin','module_policy_save',subject_type='department',subject_id='operations',overrides={'scraper':True})
    call(team,'admin','module_policy_save',subject_type='user',subject_id=team[2]['b']['id'],overrides={'scraper':False})
    assert call(team,'a','modules')['scraper']['allowed'] and not call(team,'b','modules')['scraper']['allowed']
    assert all(v['allowed'] for v in call(team,'admin','modules').values())


def test_module_grants_do_not_bypass_record_role_or_field_permissions(team):
    lid=commit(team,'a')['committed'][0]['id']
    call(team,'admin','module_policy_save',subject_type='user',subject_id=team[2]['b']['id'],overrides={'leads':True,'tools':True,'campaigns':True})
    with pytest.raises(ServiceError):call(team,'b','lead',lead_id=lid)
    with pytest.raises(ServiceError):call(team,'b','campaign_save',campaign_id='HACK',campaign_name='No',config={})
    with pytest.raises(ServiceError):call(team,'b','module_policy_save',subject_type='user',subject_id=team[2]['b']['id'],overrides={'scraper':True})
    assert all(r['user_id']==team[2]['b']['id'] for r in call(team,'b','audit'))
    call(team,'admin','module_policy_save',subject_type='user',subject_id=team[2]['a']['id'],overrides={'leads':False,'settings':False})
    for op,args in [('lead',{'lead_id':lid}),('leads',{'campaign_id':'C1'}),('sheet_create',{'campaign_id':'C1'}),('export',{'lead_ids':[lid]}),
                    ('preferences',{'key':'app_settings','value':{},'save':True})]:
        with pytest.raises(ServiceError):call(team,'a',op,**args)
    assert call(team,'a','campaigns') and call(team,'a','notifications')
    call(team,'a','preferences',key='workspace_view',value={},save=True)
    call(team,'admin','module_policy_save',subject_type='user',subject_id=team[2]['b']['id'],overrides={'validation':False})
    with pytest.raises(ServiceError,match='Validation'):commit(team,'b',44)


def test_setup_retries_short_mismatched_and_overlong_passwords(tmp_path,monkeypatch,capsys):
    import team_server
    service=Service(tmp_path/'setup.sqlite3')
    values=iter(['Admin','admin','admin@example.com'])
    monkeypatch.setattr('builtins.input',lambda _:next(values))
    passwords=iter(['short','x'*257,'Long-valid-password-1','Mismatch-password-2',PASSWORD,PASSWORD])
    monkeypatch.setattr(team_server.getpass,'getpass',lambda _:next(passwords))
    assert team_server.setup_super_user(service)
    assert service.login('admin',PASSWORD)['user']['role']=='Super User'
    output=capsys.readouterr().out
    assert '12–256' in output and 'did not match' in output and 'Traceback' not in output and PASSWORD not in output
    monkeypatch.setattr('builtins.input',lambda _:pytest.fail('Existing setup must not prompt or replace the account'))
    assert not team_server.setup_super_user(service)


def test_user_ai_cache_does_not_open_authoritative_database(tmp_path,monkeypatch):
    import app.core.cache_db as cache_db
    from app.ai.ai_service import LocalAIService
    from app.ai.qwen_runtime import QwenRuntimeManager
    forbidden=tmp_path/'do-not-create.sqlite3'; monkeypatch.setattr(cache_db,'DB_PATH',forbidden)
    a=LocalAIService(QwenRuntimeManager(),database=tmp_path/'user-a'/'cache.sqlite3')
    b=LocalAIService(QwenRuntimeManager(),database=tmp_path/'user-b'/'cache.sqlite3')
    a.knowledge.save('test','private',{'value':'A'})
    assert b.knowledge.get('test','private') is None and not forbidden.exists()


def test_one_window_modules_filters_and_scraper_transfer(qt,tmp_path,monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    from openpyxl import Workbook
    from app.constants import STANDARD_HEADERS
    from app.ui_qt.team_window import TeamWindow
    from tests.test_team_ui import DirectClient,wait_ready
    from tests.test_team_service import record
    errors=[]
    monkeypatch.setattr(QMessageBox,'warning',lambda *a:errors.append(a[-1]))
    monkeypatch.setattr(QMessageBox,'information',lambda *a:errors.append(a[-1]))
    monkeypatch.setenv('LEAD_CLIENT_DATA_DIR',str(tmp_path/'client-cache'))
    service=Service(tmp_path/'ui.sqlite3'); service.bootstrap('Admin','admin','admin@example.com',PASSWORD)
    client=DirectClient(service,service.login('admin',PASSWORD))
    client.call('campaign_save',campaign_id='C1',campaign_name='Unified campaign',config={'require_tal_match':False})
    qa=client.call('user_save',full_name='QA member',username='qa',email='qa@example.com',password=PASSWORD,role='Employee',department_id='quality')
    window=TeamWindow(client); window.show()
    try:
        wait_ready(qt,window)
        assert not errors and window.items.count('Lead Workspace')==1
        assert not set(window.items)&{'My Leads','Incoming Leads','My Worksheets','Team Leads','Department Leads'}
        for module in ('scraper','ai','settings','campaigns','reports','dashboard','tools','validation','administration','info'):
            window.navigate(module); wait_ready(qt,window)
            assert window.stack.currentWidget()==window.module_pages[module]
        assert not errors
        window.navigate('administration'); wait_ready(qt,window)
        window.access_subject.setCurrentIndex(window.access_subject.findData('quality')); wait_ready(qt,window)
        window.access_controls['scraper'].setCurrentIndex(1); window.save_module_policy(); wait_ready(qt,window)
        qa_token=service.login('qa',PASSWORD)['token']
        assert service.call(qa_token,'modules')['scraper']['allowed']
        # Stop during the asynchronous authorization step must not start a worker.
        from types import SimpleNamespace
        closed=[]; worker_calls=[]
        monkeypatch.setattr(window.scraper_page,'_ensure_worker',lambda:worker_calls.append('started'))
        window.scraper_page._cancel_preparation=False
        window.scraper_page._begin_extraction(SimpleNamespace(wb=SimpleNamespace(close=lambda:closed.append(True))))
        window.scraper_page.stop(); wait_ready(qt,window)
        assert closed==[True] and not worker_calls
        window.resize(1366,768); qt.processEvents()
        window.navigate('leads'); wait_ready(qt,window); qt.processEvents()
        assert window.leads.search.height()>=window.leads.search.minimumSizeHint().height()
        # Real worksheet ingestion, then the same grid, without opening another app.
        path=tmp_path/'ScraperResults.xlsx'; book=Workbook(); ws=book.active; ws.title='Leads'; ws.append(STANDARD_HEADERS)
        for row in [record(80),record(81),record(80)]:ws.append([row.get(h,'-') for h in STANDARD_HEADERS])
        book.save(path)
        window.import_scraper_results(str(path)); wait_ready(qt,window)
        assert window.stack.currentWidget()==window.workspace and window.controller.mode=='draft'
        assert len(window.controller.data_manager.records)==2
        assert window.controller.sheet['name'].startswith('Scraper')
        assert not client.call('leads',campaign_id='C1')['total']
        window.set_workspace_mode('all'); wait_ready(qt,window)
        assert window.controller.mode=='records' and window.stack.currentWidget()==window.workspace
        window.toggle_sidebar(); window.navigate('settings'); wait_ready(qt,window)
        assert window.current_module=='settings' and not errors
        window.settings_page.color_scheme.setCurrentText('Teal'); window.settings_page.apply(); wait_ready(qt,window)
        assert client.call('preferences')['app_settings']['color_scheme']=='Teal'
        window.grab().save(str(tmp_path/'unified-settings.png'))
    finally:
        window._closing=True; window.close(); qt.processEvents()


def test_authorized_export_preserves_scope_and_validation_fields(qt,tmp_path,monkeypatch):
    from PySide6.QtWidgets import QFileDialog,QMessageBox
    from openpyxl import load_workbook
    from app.ui_qt.team_window import TeamWindow
    from app.ui_qt.dialogs.workspace_dialogs import ExportChoiceDialog
    from tests.test_team_ui import DirectClient,wait_ready
    from tests.test_team_service import record
    errors=[]; monkeypatch.setattr(QMessageBox,'warning',lambda *a:errors.append(a[-1]))
    service=Service(tmp_path/'export.sqlite3'); service.bootstrap('Admin','admin','admin@example.com',PASSWORD)
    client=DirectClient(service,service.login('admin',PASSWORD))
    client.call('campaign_save',campaign_id='C1',campaign_name='Export',config={'require_tal_match':False})
    client.call('commit',campaign_id='C1',records=[record(23),record(24)],request_key='export-test')
    output=tmp_path/'export.xlsx'
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,**k:(str(output),'Excel (*.xlsx)'))
    def choose(dialog):dialog.scope.setCurrentIndex(dialog.scope.findData('filtered')); return 1
    monkeypatch.setattr(ExportChoiceDialog,'exec',choose)
    window=TeamWindow(client); window.show()
    try:
        wait_ready(qt,window); window.set_workspace_mode('all'); wait_ready(qt,window)
        window.leads.search.setText('Person23'); window.export_leads(); wait_ready(qt,window)
        book=load_workbook(output); assert book['Lead Data'].max_row==2 and book['Validation Details'].max_row==2
        assert not errors
    finally:window._closing=True; window.close(); qt.processEvents()


def test_quality_user_scraper_grant_deny_updates_same_window(qt,tmp_path,monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    from app.ui_qt.team_window import TeamWindow
    from tests.test_team_ui import DirectClient,wait_ready
    errors=[]; monkeypatch.setattr(QMessageBox,'warning',lambda *a:errors.append(a[-1]))
    service=Service(tmp_path/'qa-ui.sqlite3'); service.bootstrap('Admin','admin','admin@example.com',PASSWORD)
    admin=service.login('admin',PASSWORD)['token']
    user=service.call(admin,'user_save',dict(full_name='Quality user',username='qa',email='qa@example.com',password=PASSWORD,role='Employee',department_id='quality'))
    service.call(admin,'campaign_save',dict(campaign_id='C1',campaign_name='Quality campaign',config={'require_tal_match':False}))
    service.call(admin,'campaign_assign',dict(campaign_id='C1',user_ids=[user['id']]))
    client=DirectClient(service,service.login('qa',PASSWORD)); window=TeamWindow(client); window.show()
    try:
        wait_ready(qt,window); assert 'Scraper' not in window.items and not hasattr(window,'scraper_page')
        window.controller.edit_cells({(0,'First Name'):'Unsaved work'}); sid=window.controller.sheet['id']
        service.call(admin,'module_policy_save',dict(subject_type='user',subject_id=user['id'],overrides={'scraper':True,'leads':False}))
        window.refresh_access(); wait_ready(qt,window)
        assert 'Scraper' in window.items and 'Lead Workspace' not in window.items
        window.navigate('scraper'); wait_ready(qt,window); assert window.stack.currentWidget()==window.scraper_page
        assert window.controller.dirty and window.controller.data_manager.records[0]['First Name']=='Unsaved work'
        service.call(admin,'module_policy_save',dict(subject_type='user',subject_id=user['id'],overrides={'scraper':False}))
        called=[]; monkeypatch.setattr(window.scraper_page,'_ensure_worker',lambda:called.append('worker'))
        window.scraper_page.open_browser(); wait_ready(qt,window)
        assert not called and errors and 'Scraper' in errors[-1]
        window.refresh_access(); wait_ready(qt,window)
        assert 'Scraper' not in window.items and not window.scraper_page.isEnabled()
        service.call(admin,'module_policy_save',dict(subject_type='user',subject_id=user['id'],overrides={'leads':True}))
        window.refresh_access(); wait_ready(qt,window)
        assert client.call('sheet_load',sheet_id=sid)['payload']['records'][0]['First Name']=='Unsaved work'
    finally:window._closing=True; window.close(); qt.processEvents()
