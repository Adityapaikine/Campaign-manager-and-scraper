import os
import subprocess
import sys
from pathlib import Path
from app.team.mode import enable_shared_mode,shared_mode_requested


def test_all_launchers_use_one_authenticated_workspace(tmp_path):
    database=tmp_path/'lead_manager.sqlite3'
    assert shared_mode_requested(database,[],{})
    enable_shared_mode(database)
    assert shared_mode_requested(database,[],{})
    assert shared_mode_requested(database,['--research'],{'LEAD_TEAM_SERVER':'http://127.0.0.1:8765'})
    env={**os.environ,'LEAD_WORKSPACE_DATA_DIR':str(tmp_path)}
    result=subprocess.run([sys.executable,'-c','from suite_bootstrap import configure; print(configure()); print(configure())','--research'],env=env,capture_output=True,text=True,check=True)
    assert result.stdout.splitlines()==[str(tmp_path)]*2


def test_all_report_tabs_export_with_safe_values(tmp_path):
    from app.ui_qt.team_actions import write_report
    from openpyxl import load_workbook
    report={'total':1,'active_employees':1,'submitted_today_utc':0,'transfer_events':2,'mean_delivered_elapsed_hours':None,'definitions':'One Lead ID once',
            'states':{'Delivered':1},'validation':{'Valid':1},'workload':[{'full_name':'=unsafe()','assigned':1}],'departments':[],'quality_by_created_date':[]}
    output=tmp_path/'report.xlsx'; write_report(output,report)
    book=load_workbook(output)
    assert book.sheetnames==['Overview','Workflow','Validation','Workload','Departments','Quality by date']
    assert book['Workload']['A2'].data_type!='f'
