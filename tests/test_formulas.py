"""Formula integration tests; all controllers use the isolated conftest workspace."""
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import Qt, QItemSelection, QItemSelectionModel
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QDialog

from app.core.data_manager import DataManager
from app.core.formulas import function_catalog, translate_formula
from app.ui_qt.controller import ApplicationController
from app.ui_qt.dialogs.formula_dialog import FormulaDialog, DeleteColumnsDialog
from app.ui_qt.pages.leads_page import LeadsPage
from app.ui_qt.widgets.lead_details import LeadEditDialog


def sheet():
    manager = DataManager()
    manager.add_columns(["Units", "Price", "Total", "Other"])
    manager.extend([
        {"First Name": " Ada ", "Last Name": "Lovelace", "Units": "3", "Price": "20", "Country": "India"},
        {"First Name": "Grace", "Last Name": "Hopper", "Units": "4", "Price": "10", "Country": "USA"},
        {"First Name": "Katherine", "Last Name": "Johnson", "Units": "5", "Price": "30", "Country": "India"},
    ], "Formula tests")
    return manager


def controller():
    result = ApplicationController()
    result.data_manager = sheet()
    return result


@pytest.mark.parametrize("expression,expected", [
    ("=[Units]*[Price]", "60"), ("=Units+Price", "23"), ("=column43+column44", "23"),
    ("=SUM(AQ1:AQ3)", "12"), ("=SUM(AQ:AQ)", "12"), ("=AVERAGE(AQ1:AQ3)", "4"),
    ('=SUMIF(AB1:AB3,"India",AQ1:AQ3)', "8"), ('=COUNTIF(AB1:AB3,"India")', "2"),
    ('=IF([Units]>2,"Yes","No")', "Yes"), ("=IFERROR(1/0,42)", "42"),
    ('=TRIM([First Name])&" "&[Last Name]', "Ada Lovelace"),
    ('=TEXTJOIN(" ",TRUE,TRIM([First Name]),[Last Name])', "Ada Lovelace"),
    ('=LOWER([Country])', "india"), ('=LEFT([Last Name],4)', "Love"),
    ('=XLOOKUP(4,AQ1:AQ3,AR1:AR3)', "10"), ('=VLOOKUP(4,AQ1:AR3,2,FALSE)', "10"),
    ('=INDEX(AQ1:AR3,MATCH(5,AQ1:AQ3,0),2)', "30"),
    ('=TEXT(DATE(2026,10,24),"yyyy-mm-dd")', "2026-10-24"),
    ('=ROUND(PMT(0,10,100),2)', "-10"), ("=DEC2HEX(255)", "FF"),
    ("=MEDIAN(AQ1:AQ3)", "4"), ("=ROW()+COLUMN()", "46"),
    ('=IF(A4="",42,A4)', "42"), ('=IFERROR(UNKNOWN(1),42)', "#NAME?"),
    ("=1/0", "#DIV/0!"), ("=ZZA1", "#REF!"), ("=AS1+1", "#CYCLE!"),
    ("=[Missing field]", "#REF!"), ('=WEBSERVICE("https://example.com")', "#NAME?"),
    ("='Other sheet'!A1", "#REF!"), ("=[QA Status]", "#REF!"),
    ("=SUM(AQ1:AQ100001)", "#NUM!"), ("=FILTER(AQ1:AQ3,AQ1:AQ3>3)", "#SPILL!"),
    ("=SUM(FILTER(AQ1:AQ3,AQ1:AQ3>3))", "9"), ('="Text with ! and [brackets]"', "Text with ! and [brackets]"),
    ("=ROW(A2)", "2"), ("=COLUMN(AQ1)", "43"), ("=ROWS(AQ1:AQ3)", "3"),
    ("=COLUMNS(AQ1:AR3)", "2"), ("=ISREF(AQ2)", "TRUE"),
    ('=REPT("a",1000000000)', "#NUM!"), ("=MUNIT(1000000000)", "#NUM!"),
    ("=SUM(MUNIT(3))", "3"), ('=REPT("a",3)', "aaa"),
    ('=LET(amount,100,amount*1.1)', "110"), ("=SUM(MAKEARRAY(2,2,LAMBDA(r,c,r+c)))", "12"),
])
def test_calculation_families_and_errors(expression, expected):
    manager = sheet(); manager.update_cell(0, "Total", expression, allow_formulas=True)
    manager.recalculate()
    assert manager.records[0]["Total"] == expected
    assert manager.raw_value(0, "Total") == expression
    if expected.startswith("#"):
        assert manager.formula_errors[0, "Total"]


def test_user_named_column1_takes_precedence_and_brackets_support_cell_like_names():
    manager = sheet(); manager.add_columns(["column1", "column2", "A1"])
    manager.update_cell(0, "column1", "7"); manager.update_cell(0, "column2", "11"); manager.update_cell(0, "A1", "13")
    manager.update_cell(0, "Total", "=column1+column2+[A1]", allow_formulas=True); manager.recalculate()
    assert manager.records[0]["Total"] == "31"


def test_recalculation_chains_undo_and_dirty_dependent_rows():
    c = controller()
    c.edit_cells({(0, "Total"): "=[Units]*[Price]", (1, "Total"): "=AS1+1", (2, "Total"): "=AS2+1"})
    c.dirty_rows.clear()
    c.edit_cells({(0, "Units"): "10"})
    assert [r["Total"] for r in c.data_manager.records] == ["200", "201", "202"]
    assert c.dirty_rows == {0, 1, 2}
    c.undo_edit(); assert [r["Total"] for r in c.data_manager.records] == ["60", "61", "62"]
    c.redo_edit(); assert c.data_manager.records[2]["Total"] == "202"
    c.edit_cells({(0, "Total"): "=SUM(AQ1:AQ3)"})
    c.undo_edit(); assert c.data_manager.raw_value(0, "Total") == "=[Units]*[Price]"


def test_cycles_can_be_fixed_and_unrelated_formulas_calculate():
    c = controller()
    c.edit_cells({(0, "Total"): "=AS2", (1, "Total"): "=AS1", (2, "Total"): "=AS1+1", (0, "Other"): "=2+2"})
    assert [r["Total"] for r in c.data_manager.records] == ["#CYCLE!"] * 3
    assert c.data_manager.records[0]["Other"] == "4"
    c.edit_cells({(1, "Total"): "=7"})
    assert [r["Total"] for r in c.data_manager.records] == ["7", "7", "8"]


def test_column_deletion_preserves_references_and_undo_restores_formulas():
    c = controller()
    c.edit_cells({(0, "Total"): "=[Units]*[Price]", (0, "Other"): "=AS1+1", (1, "Other"): "=SUM(AQ2:AS2)"})
    c.change_columns(["Units"], delete=True)
    assert c.data_manager.records[0]["Total"] == "#REF!"
    assert c.data_manager.records[0]["Other"] == "#REF!"
    assert c.data_manager.records[1]["Other"] == "10"
    c.undo_edit()
    assert c.data_manager.custom_columns == ["Units", "Price", "Total", "Other"]
    assert c.data_manager.records[0]["Total"] == "60"
    assert c.data_manager.raw_value(0, "Other") == "=AS1+1"
    c.redo_edit(); assert "Units" not in c.data_manager.headers
    with pytest.raises(ValueError): c.change_columns(["First Name"], delete=True)


def test_add_columns_undo_and_row_deletion_rebase():
    c = controller(); c.change_columns(["New calculation"])
    c.undo_edit(); assert "New calculation" not in c.data_manager.headers
    c.redo_edit(); assert "New calculation" in c.data_manager.headers
    c.edit_cells({(2, "Total"): "=AQ2+AQ3"})
    c.data_manager.delete_indices([0])
    assert c.data_manager.raw_value(1, "Total") == "=AQ1+AQ2"
    assert c.data_manager.records[1]["Total"] == "9"


def test_external_and_imported_equals_strings_are_literal_including_undo():
    c = controller(); m = c.data_manager
    m.append_record({"Total": "=1+1"})
    assert (3, "Total") not in m.formulas
    c.edit_cells({(0, "Total"): "=1+1"}, allow_formulas=False)
    assert m.records[0]["Total"] == "=1+1" and (0, "Total") not in m.formulas
    c.edit_cells({(0, "Total"): "=5+5"})
    c.undo_edit(); assert m.records[0]["Total"] == "=1+1" and (0, "Total") not in m.formulas
    c.edit_cells({(0, "Total"): "'=6+6"})
    assert m.records[0]["Total"] == "=6+6" and (0, "Total") not in m.formulas
    c.edit_cells({(0, "Total"): "'=literal"}, allow_formulas=False)
    assert m.records[0]["Total"] == "'=literal"
    c.edit_cells({(0, "Total"): "changed"}); c.undo_edit()
    assert m.records[0]["Total"] == "'=literal" and (0, "Total") not in m.formulas


def test_workspace_checkpoint_and_export_round_trip(tmp_path):
    c = controller()
    c.edit_cells({(0, "Total"): "=[Units]*[Price]", (1, "Total"): "=AS1+2"})
    c.save_workspace()
    restored = ApplicationController()
    assert restored.data_manager.raw_value(0, "Total") == "=[Units]*[Price]"
    assert restored.data_manager.records[1]["Total"] == "62"
    checkpoint = c.save_checkpoint(reason="formula-test")
    snapshot = c.checkpoints.load(checkpoint)
    c.edit_cells({(0, "Total"): "0"}); c.recover(snapshot)
    assert c.data_manager.raw_value(0, "Total") == "=[Units]*[Price]"
    output = c.export(str(tmp_path / "formulas.xlsx"), scope="filtered", row_indices=[1])
    book = load_workbook(output)
    assert book["Lead Data"].cell(2, 45).value == "62"
    audit = book["Worksheet Formulas"]
    assert "=AS1+2" in audit.cell(2, 3).value and audit.cell(2, 3).data_type != "f"
    assert audit.cell(2, 1).value == "2"
    book.close()


def test_relative_absolute_and_context_references_in_fill():
    manager = sheet()
    expression = "=[Units]*$AR$1+ROW()+COLUMN()"
    translated = translate_formula(expression, (0, 44), (1, 44), manager.headers)
    assert translated == "=AQ2*$AR$1+ROW()+COLUMN()"
    manager.update_cell(1, "Total", translated, allow_formulas=True); manager.recalculate()
    assert manager.records[1]["Total"] == "127"


def test_appended_formula_row_is_removed_on_undo_and_recreated_on_redo():
    c = controller(); c.edit_cells({(0, "Total"): "=AS4+1"})
    c.edit_cells({(3, "Total"): "=10"})
    assert c.data_manager.records[0]["Total"] == "11"
    c.undo_edit()
    assert len(c.data_manager.records) == 3 and (3, "Total") not in c.data_manager.formulas
    assert c.data_manager.records[0]["Total"] == "1"
    c.redo_edit(); assert c.data_manager.records[0]["Total"] == "11"


def test_equals_key_begins_formula_edit_and_f2_shows_expression(qt):
    from PySide6.QtWidgets import QApplication, QLineEdit
    c = controller(); page = LeadsPage(c); page.resize(1300, 800); page.show(); page.activateWindow()
    try:
        page.table.setCurrentIndex(page.proxy.mapFromSource(page.model.index(0, 44)))
        page.table.setFocus(); qt.processEvents()
        QTest.keyClicks(page.table, "="); qt.processEvents()
        editor = QApplication.focusWidget()
        assert isinstance(editor, QLineEdit)
        QTest.keyClicks(editor, "1+2"); QTest.keyClick(editor, Qt.Key.Key_Return); qt.processEvents()
        assert c.data_manager.records[0]["Total"] == "3"
        page.table.setCurrentIndex(page.proxy.mapFromSource(page.model.index(0, 44)))
        QTest.keyClick(page.table, Qt.Key.Key_F2); qt.processEvents()
        editors = page.table.findChildren(QLineEdit)
        assert any(editor.isVisible() and editor.text() == "=1+2" for editor in editors)
        QTest.keyClick(page.table, Qt.Key.Key_Escape)
    finally: page.close()


def select_rectangle(page, top, bottom, left, right):
    selection = QItemSelection(page.proxy.index(top, left), page.proxy.index(bottom, right))
    page.table.selectionModel().select(selection, QItemSelectionModel.SelectionFlag.ClearAndSelect)


def test_formula_bar_sorted_row_edit_raw_editor_shortcuts_and_busy(qt):
    c = controller(); page = LeadsPage(c); page.resize(1320, 800); page.show()
    try:
        page.table.sortByColumn(43, Qt.SortOrder.AscendingOrder)
        page.table.setCurrentIndex(page.proxy.index(0, 44))
        assert page.formula_address.text() == "AS2"
        page.formula_entry.setText("=[Units]*[Price]"); page.apply_formula_entry()
        assert c.data_manager.records[1]["Total"] == "40"
        assert page.model.data(page.model.index(1, 44), Qt.ItemDataRole.EditRole) == "=[Units]*[Price]"
        page.formula_entry.setFocus(); page.formula_entry.setText("=AQ2"); page.formula_entry.setCursorPosition(4)
        QTest.keyClick(page.formula_entry, Qt.Key.Key_F4)
        assert page.formula_entry.text() == "=$AQ$2"
        QTest.keyClick(page.formula_entry, Qt.Key.Key_Escape)
        assert page.formula_entry.text() == "=[Units]*[Price]"
        page.set_data_busy(True); page.formula_entry.setText("=999"); page.apply_formula_entry()
        assert c.data_manager.records[1]["Total"] == "40"
        page.set_data_busy(False)
        page.show_formulas_action.trigger()
        assert page.model.data(page.model.index(1, 44)) == "=[Units]*[Price]"
        page.show_formulas_action.trigger()
        page.activateWindow(); qt.processEvents(); page.table.setFocus(); qt.processEvents()
        QTest.keyClick(page.table, Qt.Key.Key_F, Qt.KeyboardModifier.ControlModifier | Qt.KeyboardModifier.ShiftModifier)
        assert page.formula_entry.hasFocus()
    finally: page.close()


def test_fill_uses_selected_visible_rows_and_values_conversion_undo(qt):
    c = controller(); page = LeadsPage(c)
    try:
        c.edit_cells({(0, "Total"): "=[Units]*[Price]"}); page.refresh()
        page.proxy.set_column_filter(27, {"India"})
        select_rectangle(page, 0, 1, 44, 44)
        page.fill_down()
        assert c.data_manager.records[2]["Total"] == "150"
        assert c.data_manager.records[1]["Total"] == "-"
        select_rectangle(page, 0, 1, 44, 44); page.formulas_to_values()
        assert not c.data_manager.formulas
        page.undo_edit(); assert len(c.data_manager.formulas) == 2
    finally: page.close()


def test_details_edit_preserves_formula_in_an_unchanged_field(qt):
    c = controller(); c.edit_cells({(0, "Business Email ID"): '=LOWER([First Name])&"@example.com"'})
    page = LeadsPage(c)
    try:
        def edit(dialog):
            assert dialog.fields["Business Email ID"].text().startswith("=LOWER")
            dialog.fields["Job Title"].setText("Engineer")
            return QDialog.DialogCode.Accepted
        with patch.object(LeadEditDialog, "exec", edit): page.edit_lead(0)
        assert c.data_manager.raw_value(0, "Business Email ID").startswith("=LOWER")
    finally: page.close()


def test_function_search_and_explicit_column_choice(qt):
    assert len(function_catalog()) == 469
    dialog = FormulaDialog(); dialog.search.setText("XLOOKUP")
    assert dialog.functions.count() == 1 and dialog.selected_function() == "XLOOKUP"
    assert "Example:" in dialog.description.text()
    dialog.search.setText("NOT_A_FUNCTION"); assert dialog.functions.count() == 0
    deletion = DeleteColumnsDialog(None, sheet())
    assert deletion.selected_columns() == []
    deletion.columns.item(1).setCheckState(Qt.CheckState.Checked)
    assert deletion.selected_columns() == ["Price"]
    dialog.close(); deletion.close()


def test_delete_columns_ui_and_formula_view_preserve_value_sort(qt):
    c = controller(); c.edit_cells({(0, "Total"): "=8", (1, "Total"): "=100", (2, "Total"): "=3"})
    page = LeadsPage(c)
    try:
        page.table.sortByColumn(44, Qt.SortOrder.AscendingOrder)
        assert page.visible_rows() == [2, 0, 1]
        page.show_formulas_action.trigger(); assert page.visible_rows() == [2, 0, 1]
        page.table.setCurrentIndex(page.proxy.index(0, 43))
        with patch.object(DeleteColumnsDialog, "exec", return_value=QDialog.DialogCode.Accepted), patch.object(DeleteColumnsDialog, "selected_columns", return_value=["Price"]):
            page.delete_columns()
        assert "Price" not in c.data_manager.headers and page.model.columnCount() == 45
        page.undo_edit(); assert "Price" in c.data_manager.headers and page.model.columnCount() == 46
    finally: page.close()


def test_formula_layout_and_preview(qt):
    from app.ui_qt.main_window import MainWindow
    artifacts = Path(__file__).resolve().parents[1] / "TestArtifacts"
    with patch.object(MainWindow, "_offer_recovery"), patch.object(MainWindow, "refresh_qwen_status"):
        window = MainWindow()
        try:
            window.controller.data_manager = sheet()
            manager = window.controller.data_manager
            sample = [dict(row) for row in manager.records]
            for i in range(21):
                manager.append_record({**sample[i % 3], "Units": str(i+2), "Price": str(15+i*5)})
            window.controller.edit_cells({(i, "Total"): "=[Units]*[Price]" for i in range(len(manager.records))})
            window.controller.edit_cells({(i, "Other"): '=IF([Total]>=200,"Priority","Standard")' for i in range(len(manager.records))})
            page = window.leads_page
            page.custom_views["Calculations"] = {"columns": ["First Name", "Last Name", "Country", "Units", "Price", "Total", "Other"], "freeze": False}
            page._view_widths["Calculations"] = {"First Name": 140, "Last Name": 150, "Country": 120, "Units": 100, "Price": 110, "Total": 130, "Other": 170}
            page.view_selector.addItem("Calculations"); page.view_selector.setCurrentText("Calculations")
            page._details_preference = False
            window.controller.settings.appearance_mode = "Light"
            window.controller.settings.color_scheme = "Professional Blue"; window.apply_theme()
            window.resize(1440, 900); window.show(); window.navigate("leads"); page.refresh()
            page.table.setCurrentIndex(page.proxy.mapFromSource(page.model.index(0, 44)))
            qt.processEvents()
            assert page.formula_bar.height() <= 48
            assert page.formula_entry.width() > 300
            assert window.grab().save(str(artifacts / "formula_workspace.png"))
            window.resize(1120, 740); qt.processEvents()
            assert window.width() <= 1130
            assert page.table.viewport().width() > 400
            assert page.formula_entry.width() > 200
            assert window.grab().save(str(artifacts / "formula_workspace_compact.png"))
            dialog = FormulaDialog(page); dialog.search.setText("LOOKUP"); dialog.show(); qt.processEvents()
            assert dialog.grab().save(str(artifacts / "formula_function_library.png")); dialog.close()
        finally:
            window.close(); qt.processEvents()
