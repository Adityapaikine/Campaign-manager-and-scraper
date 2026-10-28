from __future__ import annotations

from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.constants import STANDARD_HEADERS
from app.utils.text_utils import excel_safe


def _format_sheet(ws, *, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for idx in range(1, ws.max_column + 1):
        header = str(ws.cell(1, idx).value or "")
        values = [header] + [str(ws.cell(r, idx).value or "") for r in range(2, min(ws.max_row, 250) + 1)]
        width = min(max((len(v) for v in values), default=12) + 2, 60)
        ws.column_dimensions[get_column_letter(idx)].width = max(width, 12)


def export_records(
    df: "pd.DataFrame",
    path: str | Path,
    *,
    validation_details: Iterable[Mapping[str, object]] | None = None,
    custom_columns: Iterable[str] = (),
    formula_details: Iterable[Mapping[str, object]] | None = None,
) -> Path:
    """Keep the standard 42 columns first, then explicitly added custom columns.

    Validation details remain on a separate sheet. Callers that do not request
    custom columns retain the original fixed-schema export contract.
    """
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)

    headers = list(STANDARD_HEADERS) + [c for c in custom_columns if c not in STANDARD_HEADERS]
    safe_df = df.reindex(columns=headers, fill_value="-").copy()
    for column in safe_df.columns:
        safe_df[column] = safe_df[column].map(excel_safe)

    safe_df.to_excel(output, index=False, engine="openpyxl", sheet_name="Lead Data")

    wb = load_workbook(output)
    ws = wb["Lead Data"]
    _format_sheet(ws)

    details = list(validation_details or [])
    if details:
        qa = wb.create_sheet("Validation Details")
        columns: list[str] = []
        for row in details:
            for key in row:
                if str(key) not in columns:
                    columns.append(str(key))
        for col_idx, key in enumerate(columns, start=1):
            qa.cell(1, col_idx, key)
        for row_idx, row in enumerate(details, start=2):
            for col_idx, key in enumerate(columns, start=1):
                qa.cell(row_idx, col_idx, excel_safe(row.get(key, "-")))
        _format_sheet(qa)

    expressions = list(formula_details or [])
    if expressions:
        audit = wb.create_sheet("Worksheet Formulas")
        keys = ["Source row", "Column", "Formula", "Calculated value"]
        audit.append(keys)
        for row in expressions:
            audit.append([excel_safe(row.get(key, "-")) for key in keys])
        _format_sheet(audit)

    wb.save(output)
    return output
