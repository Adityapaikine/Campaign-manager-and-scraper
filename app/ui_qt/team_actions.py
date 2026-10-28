import json
import uuid
from pathlib import Path
from types import SimpleNamespace
from PySide6.QtWidgets import QFileDialog,QMessageBox,QInputDialog,QListWidget,QListWidgetItem,QDialog,QVBoxLayout,QDialogButtonBox,QComboBox,QPushButton
from PySide6.QtCore import Qt
from app.ui_qt.dialogs.team_dialogs import FormDialog,RowsDialog
from app.ui_qt.dialogs.lead_import_dialog import LeadImportDialog
from app.excel.importer import read_import_tables
from app.utils.text_utils import excel_safe,normalize_header
from app.constants import STANDARD_HEADERS

class UserMappings:
    def __init__(self,prefs): self.prefs=prefs
    def signature(self,columns):return '|'.join(normalize_header(c) for c in columns)
    def get(self,name):
        mapping=self.prefs.get('mapping_'+name)
        return SimpleNamespace(mapping=mapping) if mapping else None
    def save(self,*,name,mapping,**_):self.prefs.set('mapping_'+name,mapping)

def write_excel(path,rows,sheet='Lead Data',*,include_validation=False):
    from openpyxl import Workbook
    from app.excel.exporter import _format_sheet
    book=Workbook(); ws=book.active; ws.title=sheet
    headers=list(dict.fromkeys(k for row in rows for k in row))
    ws.append([excel_safe(h) for h in headers])
    for row in rows:ws.append([excel_safe(row.get(k,'')) for k in headers])
    _format_sheet(ws)
    if include_validation and any('QA Status' in row or 'Primary Reason' in row for row in rows):
        details=book.create_sheet('Validation Details')
        fields=[h for h in ('Lead ID','CID','QA Status','Primary Reason') if h in headers]
        details.append(fields)
        for row in rows:details.append([excel_safe(row.get(h,'')) for h in fields])
        _format_sheet(details)
    book.save(path)
    return path

def write_report(path,report):
    from openpyxl import Workbook
    from app.excel.exporter import _format_sheet
    book=Workbook(); book.remove(book.active)
    sheets={'Overview':[{'Metric':k,'Value':report[k]} for k in ('total','active_employees','submitted_today_utc','transfer_events','mean_delivered_elapsed_hours','definitions')],
            'Workflow':[{'Stage':k,'Leads':v} for k,v in report['states'].items()],
            'Validation':[{'Validation':k,'Leads':v} for k,v in report['validation'].items()],
            'Workload':report['workload'],'Departments':report['departments'],'Quality by date':report['quality_by_created_date']}
    for title,rows in sheets.items():
        ws=book.create_sheet(title); headers=list(dict.fromkeys(k for row in rows for k in row))
        ws.append([excel_safe(h) for h in headers] or ['No records'])
        for row in rows:ws.append([excel_safe(row.get(k,'')) for k in headers])
        _format_sheet(ws)
    book.save(path); return path

class TeamActions:
    def recover_draft(self):
        if self.busy or not self.campaign or not self.controller.data_manager.records:return
        payload=self.controller.draft_payload()
        if QMessageBox.question(self,'Save recovery draft','Save your current permitted values and formulas as a separate private draft? Existing shared records and other sessions will not be overwritten.')!=QMessageBox.StandardButton.Yes:return
        def done(sheet):
            self.controller.dirty=False; self.save_conflict=False
            self.nav.blockSignals(True); self.nav.setCurrentRow(self.items.index('Lead Workspace')); self.nav.blockSignals(False)
            self.current_module='leads'; self.stack.setCurrentWidget(self.workspace); self.set_workspace_mode('draft',load=False)
            self.reload_sheets(sheet['id'])
        self.run(lambda:self.client.call('sheet_recovery',campaign_id=self.campaign['campaign_id'],payload=payload),done,'Saving recovery draft…')

    def copy_authorized(self,table,cut=False):
        if self.busy or not self.campaign:return
        from PySide6.QtWidgets import QApplication
        import csv,io
        proxy=table.model(); indexes=table.selectedIndexes()
        if not indexes:return
        selected={(i.row(),i.column()) for i in indexes}; visual_rows=sorted({i.row() for i in indexes})
        cols=[c for c in table.selection_columns() if any(column==c for _,column in selected)]
        rows=[proxy.mapToSource(proxy.index(r,0)).row() for r in visual_rows]
        pairs=[(v,r) for v,r in zip(visual_rows,rows) if r<len(self.controller.data_manager.records)]
        if not pairs:return
        fields=[self.controller.data_manager.headers[c] for c in cols]
        def fetch():
            if self.controller.mode=='records':return self.client.call('export',lead_ids=[self.controller.lead_rows[r]['id'] for _,r in pairs])
            sheet=self.controller.sheet
            return self.client.call('export_draft',campaign_id=self.campaign['campaign_id'],sheet_id=sheet['id'],row_ids=[sheet['payload']['row_ids'][r] for _,r in pairs])
        def copied(records):
            output=io.StringIO(); writer=csv.writer(output,delimiter='\t',lineterminator='\n'); changes={}
            for (visible,row),record in zip(pairs,records):
                writer.writerow([record.get(field,'') if (visible,col) in selected else '' for col,field in zip(cols,fields)])
                if cut:
                    for col,field in zip(cols,fields):
                        if (visible,col) in selected and field in record and self.controller.can_edit_cell(row,field):changes[row,field]=''
            QApplication.clipboard().setText(output.getvalue().rstrip('\n'))
            if changes:self.controller.edit_cells(changes,'Cut permitted cells'); self.leads.refresh()
            self.status.setText('Copied permitted fields.'+(' Non-exportable fields were omitted.' if any(any(f not in r for f in fields) for r in records) else ''))
        self.with_saved(lambda:self.run(fetch,copied,'Checking clipboard export permissions…'))

    def selection(self):
        if self.controller.mode!='records':return []
        return [{'id':self.controller.lead_rows[i]['id'],'version':self.controller.lead_rows[i]['version']} for i in self.leads.selected_rows() if i<len(self.controller.lead_rows)]

    def selected_id(self):
        selection=self.selection()
        if not selection:QMessageBox.information(self,'Select leads','Select a committed lead first.');return None
        return selection[0]['id']

    def import_leads(self):
        if self.busy:return
        if self.controller.mode!='draft' or not self.controller.sheet or self.controller.read_only_sheet:
            QMessageBox.information(self,'Import staging','Open a personal worksheet before importing. Files are staged for validation before any shared lead is committed.');return
        path,_=QFileDialog.getOpenFileName(self,'Import leads','','Excel / CSV (*.xlsx *.xlsm *.xls *.csv *.tsv)')
        if not path:return
        if Path(path).stat().st_size>50_000_000:QMessageBox.warning(self,'Import','Files are limited to 50 MB. Split the file into smaller batches.');return
        def mapped(tables):
            dialog=LeadImportDialog(self,tables,self.controller.data_manager.headers,mapping_store=UserMappings(self.controller.settings_service))
            if not dialog.exec():return
            table=dialog.selected_table(); mapping=dialog.mapping()
            if len(table)+len(self.controller.data_manager.records)>5000:QMessageBox.warning(self,'Import','A private draft sheet supports 5,000 rows. Split this import across sheets.');return
            columns=[c for c in mapping.values() if c not in STANDARD_HEADERS]
            rows=[{target:str(series[source]) for source,target in mapping.items()} for _,series in table.iterrows()]
            if QMessageBox.question(self,'Stage import',f'Add {len(rows):,} rows to this private draft? Validation and a separate commit are required to publish valid leads.')!=QMessageBox.StandardButton.Yes:return
            self.controller.data_manager.add_columns(columns); self.controller.data_manager.extend(rows,Path(path).name)
            self.controller.dirty=True; self.controller.history.clear(); self.leads.refresh(); dialog.save_mapping(); self.save(after=self.validate)
        self.run(lambda:read_import_tables(path),mapped,'Reading import file…',cancellable=True)

    def validate(self,checked=False,selected=False,indices=None):
        if self.busy or not self.campaign:return
        if self.controller.mode=='draft':
            payload=self.controller.draft_payload(); rows=payload['records']
            if not rows:return
            def done(result):
                self.validation_result=result
                self.controller.row_states={r['row']:r['status'] for r in result['results']}
                self.controller.validation_reasons={r['row']:r['messages'] for r in result['results']}
                self.leads.refresh(); self.status.setText('Validation: '+', '.join(f'{k}: {v}' for k,v in result['summary'].items()))
                self.show_validation_summary()
            self.with_saved(lambda:self.run(lambda:self.client.call('validate_draft',campaign_id=self.campaign['campaign_id'],records=rows,sheet_id=self.controller.sheet['id']),done,'Validating against the campaign…',cancellable=True))
        else:
            selection=self.selection() if selected else [{'id':r['id'],'version':r['version']} for i,r in enumerate(self.controller.lead_rows) if indices is None or i in indices]
            ids={r['id'] for r in selection}
            if selection:self.with_saved(lambda:self.run(lambda:self.client.call('validate_leads',selection=[{'id':r['id'],'version':r['version']} for r in self.controller.lead_rows if r['id'] in ids]),lambda _:self.refresh_records(),'Validating leads…'))

    def show_validation_summary(self):
        result=getattr(self,'validation_result',None)
        if result:RowsDialog(self,'Validation results',[{'Row':r['row']+1,'Category':r['category'],'Explanation':' · '.join(r['messages'])} for r in result['results']]).exec()
        elif self.controller.mode=='records':RowsDialog(self,'Validation results',[{'Lead ID':r['id'],'Category':r['validation'],'Explanation':' · '.join(r['validation_result'].get('messages',[]))} for r in self.controller.lead_rows]).exec()

    def commit_draft(self):
        if self.busy or self.controller.mode!='draft' or not self.controller.sheet or self.controller.read_only_sheet:return
        payload=self.controller.draft_payload(); rows=payload['records']; sid=self.controller.sheet['id']
        if not rows:return
        def reviewed(result):
            self.validation_result=result; valid=result['summary'].get('Valid',0)
            if not valid:self.show_validation_summary();return
            message=f"Uploaded: {len(rows):,}\n"+'\n'.join(f'{k}: {v:,}' for k,v in result['summary'].items())+f'\n\nCommit the {valid:,} valid rows? Other rows will remain in your draft. Duplicate checks run again inside the commit transaction.'
            if QMessageBox.question(self,'Commit valid records',message)!=QMessageBox.StandardButton.Yes:return
            request_key=str(uuid.uuid4())
            def done(outcome):
                accepted={r['row'] for r in outcome['committed']}
                self.controller.dirty=False
                self.load_sheet(sid)
                self.status.setText(f"Committed {len(accepted):,} leads. Invalid or duplicate rows remain in the draft.")
            self.run(lambda:self.client.call('commit',campaign_id=self.campaign['campaign_id'],records=rows,request_key=request_key,sheet_id=sid,source_name=self.controller.sheet['name']),done,'Committing valid records…')
        self.with_saved(lambda:self.run(lambda:self.client.call('validate_draft',campaign_id=self.campaign['campaign_id'],records=rows,sheet_id=sid),reviewed,'Checking the import…',cancellable=True))

    def export_leads(self):
        if self.busy or not self.campaign:return
        from app.ui_qt.dialogs.workspace_dialogs import ExportChoiceDialog
        controller=self.controller
        rows=list(range(len(controller.data_manager.records)))
        groups={'all':rows,'filtered':self.leads.visible_rows(),'selected':self.leads.selected_rows(),
                'qualified':[i for i in rows if controller.row_states.get(i) in ('Valid','Valid - Email Generated')],
                'disqualified':[i for i in rows if controller.row_states.get(i) not in ('Needs validation',None) and controller.data_manager.records[i].get('QA Status')=='Disqualified']}
        dialog=ExportChoiceDialog(self,{k:len(v) for k,v in groups.items()},default='selected' if groups['selected'] else 'qualified' if controller.settings.export_valid_only else 'filtered',shared=True)
        if not dialog.exec():return
        indices=groups[dialog.scope.currentData()]
        if not indices:return
        path,_=QFileDialog.getSaveFileName(self,'Export permitted fields',str(Path(controller.settings.default_export_folder)/'Leads.xlsx'),'Excel (*.xlsx)')
        if not path:return
        if self.controller.mode=='records':
            ids=[self.controller.lead_rows[i]['id'] for i in indices]
            fn=lambda:self.client.call('export',lead_ids=ids)
        else:
            payload=self.controller.draft_payload()
            indices=[i for i in indices if i<len(payload['records'])]
            fn=lambda:self.client.call('export_draft',campaign_id=self.campaign['campaign_id'],sheet_id=self.controller.sheet['id'],row_ids=[self.controller.sheet['payload']['row_ids'][i] for i in indices])
        self.with_saved(lambda:self.run(lambda:write_excel(path,fn(),include_validation=controller.settings.include_validation_details),lambda _:self.status.setText('Export saved with permitted fields.'),'Preparing authorized export…'))

    def export_errors(self):
        result=getattr(self,'validation_result',None)
        if not result:return
        path,_=QFileDialog.getSaveFileName(self,'Export validation errors','','Excel (*.xlsx)')
        if path:
            rows=[{'Source row':r['row']+1,'Category':r['category'],'Explanation':'; '.join(r['messages'])} for r in result['results'] if r['category']!='Valid']
            self.run(lambda:write_excel(path,rows,'Validation Errors'),lambda _:self.status.setText('Error report saved.'))

    def workflow_action(self):
        if self.busy or not self.selection():return
        def prepared():
            selection=self.selection()
            def options(result):
                states=result['transitions']
                if not states:
                    QMessageBox.information(self,'Review selected leads','There is no common next stage for this selection and your role. Select leads at the same workflow stage. Approved Operations leads can be pushed to Quality; Quality Qualified leads can be pushed to Delivery.');return
                labels={'Pending Team Lead Review':'Submit for review','Approved':'Approve selected','Quality In Progress':'Start quality review',
                    'Quality Qualified':'Mark quality qualified','Quality Disqualified':'Mark quality disqualified','Quality TBD':'Keep quality review pending',
                    'Delivery Approved':'Approve delivery','Delivered':'Mark delivered'}
                choices=[labels.get(state,state)+' → '+state for state in states]
                choice,ok=QInputDialog.getItem(self,'Review selected leads',f'Move all {len(selection)} selected lead(s) to:',choices,0,False)
                if ok:self.run(lambda:self.client.call('transition',selection=selection,state=states[choices.index(choice)]),lambda _:self.refresh_records(),'Updating workflow…')
            self.run(lambda:self.client.call('selection_actions',selection=selection),options,'Checking selected lead stages…')
        self.with_saved(prepared)

    def assign_leads(self):
        if self.busy or not self.selection():return
        def choose(users):
            dialog=QDialog(self); dialog.setWindowTitle('Assign selected leads'); root=QVBoxLayout(dialog)
            listing=QListWidget(); root.addWidget(listing)
            for user in users:
                item=QListWidgetItem(user['full_name']); item.setData(Qt.ItemDataRole.UserRole,user['id']); item.setCheckState(Qt.CheckState.Unchecked); listing.addItem(item)
            buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Ok|QDialogButtonBox.StandardButton.Cancel); buttons.accepted.connect(dialog.accept); buttons.rejected.connect(dialog.reject); root.addWidget(buttons)
            if not dialog.exec():return
            ids=[listing.item(i).data(Qt.ItemDataRole.UserRole) for i in range(listing.count()) if listing.item(i).checkState()==Qt.CheckState.Checked]
            if not ids:return
            method,ok=QInputDialog.getItem(self,'Assignment method','Method:',['Manual','Equal distribution','Round robin'],0,False)
            if ok:self.run(lambda:self.client.call('assign',selection=self.selection(),user_ids=ids,method=method),lambda _:self.refresh_records(),'Assigning leads…')
        self.with_saved(lambda:self.run(lambda:self.client.call('assignment_options',campaign_id=self.campaign['campaign_id']),choose))

    def transfer_leads(self):
        if self.busy or not self.selection():return
        from app.ui_qt.dialogs.handoff_dialog import HandoffDialog
        def prepared():
            selection=self.selection()
            def choose(preview):
                if not preview['routes']:
                    QMessageBox.information(self,'Push selected leads','No active route is configured for this department and campaign. A Super User can configure routes in Administration → Policies.');return
                dialog=HandoffDialog(self,preview)
                if not dialog.exec():return
                values=dialog.values(); destination=dialog.destination.currentData()['name']
                def done(result):
                    self.refresh_records()
                    QMessageBox.information(self,'Handoff complete',f"{result['transferred']} selected lead(s) sent to {destination}.\n\nThe receiving lead can open Incoming leads and use Assign selected. Original ownership and history are preserved.")
                self.run(lambda:self.client.call('transfer',selection=selection,**values),done,'Sending selected leads…')
            self.run(lambda:self.client.call('selection_actions',selection=selection),choose,'Checking handoff readiness…')
        self.with_saved(prepared)

    def return_rework(self):
        lid=self.selected_id()
        if not lid or self.busy:return
        def choose(options):
            dialog=FormDialog(self,'Return for correction',{'reason_id':('Rejection reason',None,[(r['name'],r['id']) for r in options['reasons']]),'comment':('Comment','',None)})
            if dialog.exec():self.run(lambda:self.client.call('rework',selection=self.selection(),**dialog.values()),lambda _:self.refresh_records(),'Returning for rework…')
        self.with_saved(lambda:self.run(lambda:self.client.call('workflow_options',lead_id=lid),choose))

    def show_history(self):
        lid=self.selected_id()
        if lid:self.run(lambda:self.client.call('history',lead_id=lid),lambda rows:RowsDialog(self,'Lead change history',rows).exec())
    def transfer_history(self):
        lid=self.selected_id()
        if lid:self.run(lambda:self.client.call('transfer_history',lead_id=lid),lambda rows:RowsDialog(self,'Transfer history',rows).exec())
    def comments(self):
        lid=self.selected_id()
        if not lid:return
        body,ok=QInputDialog.getMultiLineText(self,'Lead comments','Add a comment (leave blank to view history):')
        if ok:self.run(lambda:self.client.call('comments',lead_id=lid,body=body if body.strip() else None),lambda rows:RowsDialog(self,'Lead comments',rows).exec())

    def archive_selected(self):self._archive(False)
    def restore_selected(self):self._archive(True)
    def _archive(self,restore):
        if not self.selection() or self.busy:return
        if QMessageBox.question(self,'Restore leads' if restore else 'Archive leads',f"{'Restore' if restore else 'Archive'} {len(self.selection())} selected leads? History will be preserved.")!=QMessageBox.StandardButton.Yes:return
        self.run(lambda:self.client.call('archive',selection=self.selection(),restore=restore),lambda _:self.refresh_records())

    def assign_historical(self):
        if not self.selection() or self.busy:return
        def choose(directory):
            users=[u for u in directory['users'] if u['active'] and u['department_id']]
            dialog=FormDialog(self,'Assign historical ownership',{'owner_id':('Original owner',None,[(u['full_name'],u['id']) for u in users])})
            if dialog.exec():self.run(lambda:self.client.call('historical_assignment',selection=self.selection(),**dialog.values()),lambda _:self.refresh_records())
        self.run(lambda:self.client.call('directory'),choose)

    def find_replace(self):
        fields=[h for h in self.controller.data_manager.headers if self.controller.can_edit_cell(0,h)]
        dialog=FormDialog(self,'Find and replace in the current selection / page',{'field':('Column',None,[(h,h) for h in fields]),'find':('Find text','',None),'replacement':('Replace with','',None)})
        if not dialog.exec():return
        values=dialog.values(); field=values['field']; find=values['find']
        if not find or not field:return
        rows=self.leads.selected_rows() or self.leads.visible_rows(); edits={}
        for row in rows:
            if row<len(self.controller.data_manager.records):
                original=self.controller.data_manager.raw_value(row,field)
                if find in original:edits[row,field]=original.replace(find,values['replacement'])
        if edits and QMessageBox.question(self,'Replace values',f'Replace text in {len(edits)} permitted cells?')==QMessageBox.StandardButton.Yes:
            self.controller.edit_cells(edits,'Find and replace'); self.leads.refresh()

    def save_filter_view(self):
        name,ok=QInputDialog.getText(self,'Save filter view','View name:')
        if ok and name.strip():
            headers=self.controller.data_manager.headers
            views=dict(self.controller.settings_service.get('filter_views',{})); views[name.strip()]={'search':self.leads.search.text(),'status':self.leads.status.currentText(),
              'fields':{headers[k]:list(v) for k,v in self.leads.proxy.column_filters.items() if k<len(headers)},
              'conditions':{headers[k]:v for k,v in self.leads.proxy.column_conditions.items() if k<len(headers)}}
            self.controller.settings_service.set('filter_views',views); self.status.setText('Filter view saved for your account.')

    def load_filter_view(self):
        views=self.controller.settings_service.get('filter_views',{})
        if not views:return
        name,ok=QInputDialog.getItem(self,'Saved filter views','View:',list(views),0,False)
        if ok:
            view=views[name]; self.leads.clear_filters(); self.leads.search.setText(view['search']); self.leads.status.setCurrentText(view['status'])
            headers=self.controller.data_manager.headers
            for field,values in view.get('fields',{}).items():
                if field in headers:self.leads.proxy.set_column_filter(headers.index(field),set(values))
            for field,rule in view.get('conditions',{}).items():
                if field in headers:self.leads.proxy.set_condition(headers.index(field),rule)
            for col,values in view.get('columns',{}).items():
                if 0<=int(col)<len(headers):self.leads.proxy.set_column_filter(int(col),set(values))

    def show_notifications(self,event=None):
        event=event if isinstance(event,str) else None
        def show(result):
            dialog=RowsDialog(self,'Notifications',result['rows'])
            events=QComboBox(); events.addItem('All event types',None)
            for name in result['events']:events.addItem(name,name)
            events.setCurrentIndex(max(0,events.findData(event))); dialog.layout().insertWidget(0,events)
            def changed():dialog.accept(); self.show_notifications(events.currentData())
            events.currentIndexChanged.connect(changed)
            mark=QPushButton('Mark displayed notifications as read'); dialog.layout().insertWidget(1,mark)
            def read():
                dialog.accept(); self.run(lambda:self.client.call('notifications',read_ids=[r['id'] for r in result['rows']]),lambda r:self.notification_button.setText(f"Notifications ({r['unread']})"))
            mark.clicked.connect(read)
            dialog.table.itemDoubleClicked.connect(lambda item:self._notification_open(result['rows'][item.row()],dialog))
            dialog.exec()
        self.run(lambda:self.client.call('notifications',event=event),show)
    def _notification_open(self,row,dialog):
        dialog.accept()
        if row.get('lead_id'):
            self.run(lambda:self.client.call('lead',lead_id=row['lead_id']),lambda lead:RowsDialog(self,'Lead '+lead['id'],[{'Field':k,'Value':v} for k,v in {**{k:v for k,v in lead.items() if k not in ('record','field_access','validation_result')},**lead['record']}.items()]).exec())
        elif row.get('campaign_id'):self.select_campaign_id(row['campaign_id'])

    def change_password(self):
        dialog=FormDialog(self,'Change password',{'current_password':('Current password','',None),'new_password':('New password (12+ characters)','',None)})
        if dialog.exec():self.run(lambda:self.client.call('change_password',**dialog.values()),lambda _:self._finish_close())
