from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.ai.ai_service import LocalAIService
from app.ai.structured_output import clamp_confidence
from app.core.campaign_models import CampaignProfile, CampaignRule, RuleSource
from app.core.document_reader import DocumentReader


class QwenCampaignReader:
    """AI second-pass reader over specification-like campaign files.

    Deterministic extraction remains primary. This reader is intentionally opt-in
    and only receives small specification files/chunks, never TAL/suppression lists.
    """

    def __init__(self, ai: LocalAIService) -> None:
        self.ai = ai
        self.documents = DocumentReader()

    def enhance_profile(self, profile: CampaignProfile) -> list[CampaignRule]:
        texts: list[dict[str, str]] = []
        for info in profile.files:
            if info.role != "specification":
                continue
            path = Path(info.absolute_path)
            if not path.exists():
                continue
            text = self._read_text(path)
            if text.strip():
                texts.append({"file": info.relative_path, "text": text[:8000]})
        if not texts:
            return []

        existing = [
            {"field": r.field, "operator": r.operator, "value": r.value, "status": r.status}
            for r in profile.rules
        ]
        result = self.ai.run_json(
            "campaign_requirements",
            {
                "campaign_id": profile.campaign_id,
                "campaign_name": profile.campaign_name,
                "existing_rules": existing,
                "documents": texts,
            },
        )
        return self._to_rules(result, texts)

    def _to_rules(self, result: dict[str, Any], texts: list[dict[str, str]]) -> list[CampaignRule]:
        rules: list[CampaignRule] = []
        for item in result.get("rules", []) if isinstance(result.get("rules"), list) else []:
            if not isinstance(item, dict):
                continue
            field = str(item.get("field") or "").strip()
            if not field:
                continue
            confidence = clamp_confidence(item.get("confidence"))
            review = bool(item.get("review_required", False)) or confidence < self.ai.config.minimum_confidence
            excerpt = str(item.get("source_excerpt") or "").strip()
            source_file = str(item.get("source_file") or "").strip()
            if not source_file and texts:
                source_file = texts[0]["file"]
            rules.append(
                CampaignRule(
                    category=str(item.get("category") or "AI Extracted"),
                    field=field,
                    operator=str(item.get("operator") or "equals"),
                    value=item.get("value"),
                    confidence=confidence,
                    status="Needs Review" if review else "Extracted",
                    notes=f"Qwen: {str(item.get('reason') or '').strip()}".strip(),
                    sources=[RuleSource(source_file or "Qwen", "AI extraction", excerpt)],
                )
            )
        return rules

    def _read_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".docx", ".doc"}:
            return self.documents.read(path).text
        if suffix in {".xlsx", ".xlsm"}:
            wb = load_workbook(path, read_only=True, data_only=True)
            parts: list[str] = []
            for ws in wb.worksheets[:8]:
                parts.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if values:
                        parts.append(" | ".join(values))
                    if sum(len(x) for x in parts) > 10000:
                        break
            return "\n".join(parts)
        if suffix == ".xls":
            try:
                import xlrd
                book = xlrd.open_workbook(str(path), on_demand=True)
                parts = []
                for sheet in book.sheets()[:8]:
                    parts.append(f"[Sheet: {sheet.name}]")
                    for r in range(min(sheet.nrows, 500)):
                        values = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols) if str(sheet.cell_value(r, c)).strip()]
                        if values:
                            parts.append(" | ".join(values))
                return "\n".join(parts)
            except Exception:
                return ""
        if suffix in {".csv", ".tsv", ".txt"}:
            for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
                try:
                    raw = path.read_text(encoding=encoding)
                    return raw[:12000]
                except UnicodeDecodeError:
                    continue
        return ""
