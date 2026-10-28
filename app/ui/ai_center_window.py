from __future__ import annotations

import json
import threading
from tkinter import filedialog, messagebox

import customtkinter as ctk

from app.ai.campaign_intelligence import AICampaignIntelligence
from app.ai.enrichment_agent import AIEnrichmentAgent
from app.ai.normalization_agent import AINormalizationAgent
from app.ai.qa_assistant import AIQAAssistant
from app.ai.research_assistant import AIResearchAssistant
from app.ai.rule_builder import AIRuleBuilder
from app.core.activity_history import record_activity
from app.core.audit_service import record_audit_change
from app.core.campaign_models import CampaignRule
from app.core.validation_preview import ValidationPreviewService
from app.services.search_provider import create_search_provider


class AICenterWindow(ctk.CTkToplevel):
    """Lightweight UI facade over optional local-Qwen features.

    No model is loaded by Python. Every AI action runs on a worker thread and the
    window remains responsive while the local server is working.
    """

    def __init__(self, master):
        super().__init__(master)
        self.title("Local Qwen AI Center")
        self.geometry("1180x800")
        self.minsize(900, 650)
        self.transient(master)
        self.grid_rowconfigure(1, weight=1)
        self.grid_columnconfigure(0, weight=1)
        self.master_app = master
        self.ai = master.ai_service
        self._last_normalization: dict[str, dict] = {}
        self._last_enrichment: dict[str, dict] = {}
        self._last_enrichment_row: int | None = None
        self._last_management_report: dict | None = None

        header = ctk.CTkFrame(self)
        header.grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 6))
        header.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(header, text="Local Qwen AI Center", font=ctk.CTkFont(size=22, weight="bold")).grid(
            row=0, column=0, sticky="w", padx=14, pady=10
        )
        self.status = ctk.CTkLabel(header, text="Ready")
        self.status.grid(row=0, column=1, sticky="w", padx=12)
        ctk.CTkButton(header, text="Qwen Settings", width=115, command=master.open_settings).grid(row=0, column=2, padx=(6, 14))

        self.tabs = ctk.CTkTabview(self)
        self.tabs.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0, 12))
        self.campaign_tab = self.tabs.add("Campaign AI")
        self.data_tab = self.tabs.add("Data AI")
        self.qa_tab = self.tabs.add("QA Assistant")
        self.intelligence_tab = self.tabs.add("Campaign Intelligence")
        self.knowledge_tab = self.tabs.add("Knowledge & Cache")
        for tab in (self.campaign_tab, self.data_tab, self.qa_tab, self.intelligence_tab, self.knowledge_tab):
            tab.grid_rowconfigure(0, weight=1)
            tab.grid_columnconfigure(0, weight=1)

        self._build_campaign_tab()
        self._build_data_tab()
        self._build_qa_tab()
        self._build_intelligence_tab()
        self._build_knowledge_tab()

    def _panel(self, parent):
        frame = ctk.CTkScrollableFrame(parent)
        frame.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        frame.grid_columnconfigure(0, weight=1)
        return frame

    def _output_box(self, parent, height=260):
        box = ctk.CTkTextbox(parent, height=height, wrap="word")
        box.pack(fill="both", expand=True, padx=12, pady=(6, 12))
        return box

    @staticmethod
    def _set_box(box, value) -> None:
        box.configure(state="normal")
        box.delete("1.0", "end")
        if isinstance(value, (dict, list)):
            text = json.dumps(value, indent=2, ensure_ascii=False, default=str)
        else:
            text = str(value)
        box.insert("1.0", text)
        box.configure(state="disabled")

    def _run_async(self, label: str, worker, on_done=None) -> None:
        self.status.configure(text=label)

        def job():
            try:
                result = worker()
                self.after(0, lambda: self._finish_async("Ready", result, on_done))
            except Exception as exc:
                self.after(0, lambda: self._fail_async(str(exc)))

        threading.Thread(target=job, daemon=True, name=f"ai-{label[:20]}").start()

    def _finish_async(self, status: str, result, on_done) -> None:
        self.status.configure(text=status)
        if on_done:
            on_done(result)

    def _fail_async(self, error: str) -> None:
        self.status.configure(text="AI action failed")
        messagebox.showerror("Local Qwen", error, parent=self)

    # ---------------- Campaign AI ----------------
    def _build_campaign_tab(self) -> None:
        page = self._panel(self.campaign_tab)
        card = ctk.CTkFrame(page)
        card.pack(fill="x", padx=8, pady=8)
        ctk.CTkLabel(card, text="AI Rule Builder", font=ctk.CTkFont(size=17, weight="bold")).pack(anchor="w", padx=12, pady=(12, 4))
        ctk.CTkLabel(card, text="Describe campaign targeting in natural language. Suggested rules are reviewed before activation.", wraplength=1000, justify="left").pack(anchor="w", padx=12)
        self.rule_input = ctk.CTkTextbox(card, height=110)
        self.rule_input.pack(fill="x", padx=12, pady=8)
        ctk.CTkButton(card, text="Generate Suggested Rules", command=self._generate_rules).pack(anchor="w", padx=12, pady=(0, 10))
        self.rule_output = self._output_box(card, 220)

        preview = ctk.CTkFrame(page)
        preview.pack(fill="x", padx=8, pady=8)
        ctk.CTkLabel(preview, text="Validation Preview", font=ctk.CTkFont(size=17, weight="bold")).pack(anchor="w", padx=12, pady=(12, 4))
        ctk.CTkLabel(preview, text="Simulate the currently active campaign rules against the working data without changing any row.", wraplength=1000, justify="left").pack(anchor="w", padx=12)
        ctk.CTkButton(preview, text="Preview Current Rules", command=self._validation_preview).pack(anchor="w", padx=12, pady=8)
        self.preview_output = self._output_box(preview, 200)

    def _generate_rules(self) -> None:
        instruction = self.rule_input.get("1.0", "end").strip()
        if not instruction:
            messagebox.showinfo("AI Rule Builder", "Enter a campaign instruction first.", parent=self)
            return
        builder = AIRuleBuilder(self.ai)

        def done(rules: list[CampaignRule]):
            display = [rule.to_dict() for rule in rules]
            self._set_box(self.rule_output, display)
            profile = self.master_app.active_campaign_profile
            if profile and rules:
                add = messagebox.askyesno(
                    "AI Rule Builder",
                    f"Generated {len(rules)} suggested rule(s). Add them to the active campaign profile for review?",
                    parent=self,
                )
                if add:
                    profile.rules.extend(rules)
                    self.master_app.campaign_store.save(profile)
                    self.master_app.activate_campaign_profile(profile, silent=True)
                    record_activity("ai_rule_builder", f"Added {len(rules)} AI suggested rule(s)", {"campaign_id": profile.campaign_id})

        self._run_async("Building campaign rules...", lambda: builder.build(instruction), done)

    def _validation_preview(self) -> None:
        if not self.master_app.data_manager.records:
            messagebox.showinfo("Validation Preview", "Load lead data first.", parent=self)
            return
        service = ValidationPreviewService()
        snapshot = [dict(row) for row in self.master_app.data_manager.records]
        self._run_async(
            "Simulating campaign validation...",
            lambda: service.simulate(snapshot, self.master_app.validator).to_metrics(),
            lambda result: self._set_box(self.preview_output, result),
        )

    # ---------------- Data AI ----------------
    def _build_data_tab(self) -> None:
        page = self._panel(self.data_tab)
        normalize = ctk.CTkFrame(page)
        normalize.pack(fill="x", padx=8, pady=8)
        ctk.CTkLabel(normalize, text="Smart Normalization", font=ctk.CTkFont(size=17, weight="bold")).pack(anchor="w", padx=12, pady=(12, 4))
        controls = ctk.CTkFrame(normalize, fg_color="transparent")
        controls.pack(fill="x", padx=12, pady=5)
        self.normalize_field = ctk.StringVar(value="Job Title")
        ctk.CTkOptionMenu(controls, variable=self.normalize_field, values=["Job Title", "Job Level", "Job Function", "Industry", "Company Name", "Country"], width=180).pack(side="left")
        ctk.CTkButton(controls, text="Normalize Unique Values", command=self._normalize_unique).pack(side="left", padx=8)
        ctk.CTkButton(controls, text="Save High-Confidence to Knowledge", command=self._save_normalization_knowledge).pack(side="left", padx=8)
        ctk.CTkButton(controls, text="Bulk Apply Title Level/Function", command=self._bulk_apply_title_mappings).pack(side="left", padx=8)
        self.normalize_output = self._output_box(normalize, 250)

        enrich = ctk.CTkFrame(page)
        enrich.pack(fill="x", padx=8, pady=8)
        ctk.CTkLabel(enrich, text="Selected Lead Enrichment", font=ctk.CTkFont(size=17, weight="bold")).pack(anchor="w", padx=12, pady=(12, 4))
        ctk.CTkLabel(enrich, text="Qwen infers only from fields already present in the selected row. Suggestions are not written until you choose Apply.", wraplength=1000, justify="left").pack(anchor="w", padx=12)
        enrich_actions = ctk.CTkFrame(enrich, fg_color="transparent")
        enrich_actions.pack(fill="x", padx=12, pady=8)
        ctk.CTkButton(enrich_actions, text="Suggest Missing Fields", command=self._suggest_enrichment).pack(side="left", padx=(0, 6))
        ctk.CTkButton(enrich_actions, text="Research Company (optional web)", command=self._research_company).pack(side="left", padx=6)
        ctk.CTkButton(enrich_actions, text="Apply High-Confidence Suggestions", command=self._apply_enrichment).pack(side="left", padx=6)
        self.enrichment_output = self._output_box(enrich, 260)

    def _normalize_unique(self) -> None:
        field = self.normalize_field.get()
        values = [row.get(field, "-") for row in self.master_app.data_manager.records]
        agent = AINormalizationAgent(self.ai)

        def done(result):
            self._last_normalization = result
            self._set_box(self.normalize_output, result)

        self._run_async(f"Normalizing unique {field} values...", lambda: agent.normalize_unique(field, values), done)

    def _save_normalization_knowledge(self) -> None:
        if not self._last_normalization:
            messagebox.showinfo("Local Knowledge", "Run Smart Normalization first.", parent=self)
            return
        field = self.normalize_field.get()
        threshold = self.ai.config.minimum_confidence
        saved = 0
        for raw, result in self._last_normalization.items():
            confidence = float(result.get("confidence") or 0)
            if confidence < threshold:
                continue
            self.ai.knowledge.save(f"normalize:{field.casefold()}", raw, result, confidence=confidence, source="AI reviewed by user")
            saved += 1
        messagebox.showinfo("Local Knowledge", f"Saved {saved} high-confidence mapping(s).", parent=self)
        self._refresh_knowledge()

    def _bulk_apply_title_mappings(self) -> None:
        if self.normalize_field.get() != "Job Title" or not self._last_normalization:
            messagebox.showinfo("Bulk QA", "Run Smart Normalization for Job Title first.", parent=self)
            return
        threshold = self.ai.config.minimum_confidence
        lookup = {" ".join(raw.casefold().split()): item for raw, item in self._last_normalization.items()}
        changes = 0
        affected = 0
        for row_index, record in enumerate(self.master_app.data_manager.records):
            raw = str(record.get("Job Title") or "").strip()
            item = lookup.get(" ".join(raw.casefold().split()))
            if not item or float(item.get("confidence") or 0) < threshold:
                continue
            row_changed = False
            for field, key in (("Job Level", "job_level"), ("Job Function", "job_function")):
                value = str(item.get(key) or "-").strip()
                old = str(record.get(field, "-")).strip() or "-"
                if value in {"", "-"} or old not in {"", "-"}:
                    continue
                record[field] = value
                record_audit_change(
                    record_key=str(row_index + 1), field_name=field, old_value=old, new_value=value,
                    reason=f"Bulk QA mapping from Job Title '{raw}'", confidence=float(item.get("confidence") or 0),
                )
                self.ai.evidence.add(
                    record_key=str(row_index + 1), field_name=field, ai_value=value,
                    confidence=float(item.get("confidence") or 0), reason=str(item.get("reason") or "Job title normalization"),
                    sources={"Job Title": raw}, prompt_version="normalize_value:1.0",
                )
                changes += 1
                row_changed = True
            if row_changed:
                affected += 1
        if changes:
            self.master_app._mark_dirty()
            self.master_app._invalidate_validation()
            self.master_app.refresh_grid()
        messagebox.showinfo("Bulk QA", f"Applied {changes} field update(s) across {affected} row(s). Existing populated values were preserved.", parent=self)

    def _selected_row(self):
        row_id = self.master_app.grid.active_row_index
        if row_id is None or not (0 <= row_id < len(self.master_app.data_manager.records)):
            return None, None
        return row_id, self.master_app.data_manager.records[row_id]

    def _suggest_enrichment(self) -> None:
        row_id, record = self._selected_row()
        if record is None:
            messagebox.showinfo("AI Enrichment", "Select a lead row first.", parent=self)
            return
        agent = AIEnrichmentAgent(self.ai)

        def done(result):
            self._last_enrichment = result
            self._last_enrichment_row = row_id
            self._set_box(self.enrichment_output, result or {"message": "No evidence-backed suggestions were produced."})

        self._run_async("Generating evidence-bound suggestions...", lambda: agent.suggest(dict(record)), done)

    def _research_company(self) -> None:
        row_id, record = self._selected_row()
        if record is None:
            messagebox.showinfo("AI Research", "Select a lead row first.", parent=self)
            return
        provider = create_search_provider(
            self.master_app.search_provider_name,
            timeout_seconds=self.master_app.email_search_timeout,
            retries=self.master_app.email_search_retries,
        )
        if not provider.configured:
            messagebox.showinfo("AI Research", "No web search provider is configured. Configure one in .env / Email Discovery settings.", parent=self)
            return
        assistant = AIResearchAssistant(self.ai, provider)

        def done(result):
            suggestions = result.get("suggestions") if isinstance(result, dict) else None
            if isinstance(suggestions, dict):
                self._last_enrichment = suggestions
                self._last_enrichment_row = row_id
            self._set_box(self.enrichment_output, result)

        self._run_async("Researching company evidence...", lambda: assistant.research_company(dict(record)), done)

    def _apply_enrichment(self) -> None:
        if self._last_enrichment_row is None or not self._last_enrichment:
            messagebox.showinfo("AI Enrichment", "Generate suggestions first.", parent=self)
            return
        row_id = self._last_enrichment_row
        if not (0 <= row_id < len(self.master_app.data_manager.records)):
            return
        record = self.master_app.data_manager.records[row_id]
        threshold = self.ai.config.minimum_confidence
        changes = 0
        for field, item in self._last_enrichment.items():
            if not isinstance(item, dict):
                continue
            confidence = float(item.get("confidence") or 0)
            value = str(item.get("value") or "-").strip()
            old = str(record.get(field, "-")).strip() or "-"
            if confidence < threshold or value == "-" or old not in {"", "-"}:
                continue
            record[field] = value
            record_audit_change(
                record_key=str(row_id + 1), field_name=field, old_value=old, new_value=value,
                reason=f"Local Qwen suggestion: {item.get('reason', '')}", confidence=confidence,
            )
            self.ai.evidence.add(
                record_key=str(row_id + 1), field_name=field, ai_value=value, confidence=confidence,
                reason=str(item.get("reason") or ""), sources=item.get("evidence"), prompt_version="enrich_record:1.0",
            )
            changes += 1
        if changes:
            self.master_app._mark_dirty()
            self.master_app._invalidate_validation()
            self.master_app.refresh_grid()
        messagebox.showinfo("AI Enrichment", f"Applied {changes} high-confidence missing-field suggestion(s).", parent=self)

    # ---------------- QA ----------------
    def _build_qa_tab(self) -> None:
        page = self._panel(self.qa_tab)
        card = ctk.CTkFrame(page)
        card.pack(fill="x", padx=8, pady=8)
        ctk.CTkLabel(card, text="Selected Lead QA Assistant", font=ctk.CTkFont(size=17, weight="bold")).pack(anchor="w", padx=12, pady=(12, 4))
        actions = ctk.CTkFrame(card, fg_color="transparent")
        actions.pack(fill="x", padx=12, pady=8)
        ctk.CTkButton(actions, text="AI QA Review", command=self._qa_review).pack(side="left", padx=(0, 6))
        ctk.CTkButton(actions, text="Explain Validation", command=self._explain_validation).pack(side="left", padx=6)
        ctk.CTkButton(actions, text="AI Review Manual Rows", command=self._review_manual_rows).pack(side="left", padx=6)
        ctk.CTkButton(actions, text="Compare 2 Selected Rows", command=self._duplicate_review).pack(side="left", padx=6)
        ctk.CTkButton(actions, text="Review Duplicate Groups", command=self._review_duplicate_groups).pack(side="left", padx=6)
        self.qa_output = self._output_box(card, 460)

    def _qa_review(self) -> None:
        row_id, record = self._selected_row()
        if record is None:
            messagebox.showinfo("QA Assistant", "Select a lead row first.", parent=self)
            return
        assistant = AIQAAssistant(self.ai)
        profile = self.master_app.active_campaign_profile
        rules = [r.to_dict() for r in profile.rules if r.status in {"Confirmed", "Extracted"}] if profile else []
        self._run_async(
            "Reviewing selected lead...",
            lambda: assistant.review(dict(record), status=self.master_app.row_states.get(row_id), reasons=self.master_app.validation_reasons.get(row_id, []), campaign_rules=rules),
            lambda result: self._set_box(self.qa_output, result),
        )

    def _explain_validation(self) -> None:
        row_id, record = self._selected_row()
        if record is None:
            messagebox.showinfo("Validation Explanation", "Select a lead row first.", parent=self)
            return
        assistant = AIQAAssistant(self.ai)
        self._run_async(
            "Explaining validation...",
            lambda: assistant.explain_validation(dict(record), status=self.master_app.row_states.get(row_id), reasons=self.master_app.validation_reasons.get(row_id, [])),
            lambda result: self._set_box(self.qa_output, result),
        )

    def _duplicate_review(self) -> None:
        indices = self.master_app.grid.selected_indices()
        if len(indices) != 2:
            messagebox.showinfo("Duplicate Review", "Select exactly two worksheet rows.", parent=self)
            return
        left = dict(self.master_app.data_manager.records[indices[0]])
        right = dict(self.master_app.data_manager.records[indices[1]])
        assistant = AIQAAssistant(self.ai)
        self._run_async("Comparing possible duplicates...", lambda: assistant.compare_duplicates(left, right), lambda result: self._set_box(self.qa_output, result))

    def _review_manual_rows(self) -> None:
        candidates = [
            index for index, status in self.master_app.row_states.items()
            if status in {"Manual Review", "Warning"}
        ][:25]
        if not candidates:
            messagebox.showinfo("Hybrid Validation", "No Manual Review/Warning rows are available. Run VALIDATE first.", parent=self)
            return
        assistant = AIQAAssistant(self.ai)
        profile = self.master_app.active_campaign_profile
        rules = [r.to_dict() for r in profile.rules if r.status in {"Confirmed", "Extracted"}] if profile else []

        def worker():
            results = []
            for index in candidates:
                record = dict(self.master_app.data_manager.records[index])
                review = assistant.review(
                    record, status=self.master_app.row_states.get(index),
                    reasons=self.master_app.validation_reasons.get(index, []), campaign_rules=rules,
                )
                results.append({"row": index + 1, "deterministic_status": self.master_app.row_states.get(index), "ai_review": review})
            return results

        self._run_async("AI reviewing manual rows...", worker, lambda result: self._set_box(self.qa_output, result))

    def _review_duplicate_groups(self) -> None:
        pairs: set[tuple[int, int]] = set()
        for row_id, result in self.master_app.validation_results.items():
            for other in getattr(result, "duplicate_rows", ()):
                a, b = sorted((int(row_id), int(other)))
                if a != b:
                    pairs.add((a, b))
                if len(pairs) >= 20:
                    break
            if len(pairs) >= 20:
                break
        if not pairs:
            messagebox.showinfo("Duplicate Review", "No duplicate groups are available. Run VALIDATE first.", parent=self)
            return
        assistant = AIQAAssistant(self.ai)

        def worker():
            output = []
            for a, b in sorted(pairs):
                if a >= len(self.master_app.data_manager.records) or b >= len(self.master_app.data_manager.records):
                    continue
                review = assistant.compare_duplicates(
                    dict(self.master_app.data_manager.records[a]),
                    dict(self.master_app.data_manager.records[b]),
                )
                output.append({"rows": [a + 1, b + 1], "ai_review": review})
            return output

        self._run_async("Reviewing duplicate groups...", worker, lambda result: self._set_box(self.qa_output, result))

    # ---------------- Intelligence ----------------
    def _build_intelligence_tab(self) -> None:
        page = self._panel(self.intelligence_tab)
        card = ctk.CTkFrame(page)
        card.pack(fill="x", padx=8, pady=8)
        ctk.CTkLabel(card, text="Campaign Intelligence & Management Reporting", font=ctk.CTkFont(size=17, weight="bold")).pack(anchor="w", padx=12, pady=(12, 4))
        actions = ctk.CTkFrame(card, fg_color="transparent")
        actions.pack(fill="x", padx=12, pady=8)
        ctk.CTkButton(actions, text="Generate Campaign Insights", command=self._campaign_insights).pack(side="left", padx=(0, 6))
        ctk.CTkButton(actions, text="Generate Management Report", command=self._management_report).pack(side="left", padx=6)
        ctk.CTkButton(actions, text="Save Report to DOCX", command=self._save_management_report).pack(side="left", padx=6)
        self.intelligence_output = self._output_box(card, 500)

    def _metrics(self) -> dict:
        statuses: dict[str, int] = {}
        for state in self.master_app.row_states.values():
            statuses[state] = statuses.get(state, 0) + 1
        reasons: dict[str, int] = {}
        for values in self.master_app.validation_reasons.values():
            for reason in values:
                reasons[reason] = reasons.get(reason, 0) + 1
        profile = self.master_app.active_campaign_profile
        campaign = {
            "campaign_id": profile.campaign_id,
            "campaign_name": profile.campaign_name,
            "rules": [r.to_dict() for r in profile.rules if r.status in {"Confirmed", "Extracted"}],
        } if profile else None
        total = len(self.master_app.data_manager.records)
        email_present = sum(1 for row in self.master_app.data_manager.records if str(row.get("Business Email ID", "-")).strip() not in {"", "-"})
        return {
            "campaign": campaign,
            "total_records": total,
            "status_counts": statuses,
            "top_validation_reasons": dict(sorted(reasons.items(), key=lambda kv: kv[1], reverse=True)[:20]),
            "email_coverage_percent": round((email_present / total * 100), 1) if total else 0,
        }

    def _campaign_insights(self) -> None:
        intelligence = AICampaignIntelligence(self.ai)
        self._run_async("Analyzing campaign metrics...", lambda: intelligence.insights(self._metrics()), lambda result: self._set_box(self.intelligence_output, result))

    def _management_report(self) -> None:
        intelligence = AICampaignIntelligence(self.ai)

        def done(result):
            self._last_management_report = result
            self._set_box(self.intelligence_output, result)

        self._run_async("Writing management report...", lambda: intelligence.management_report(self._metrics()), done)

    def _save_management_report(self) -> None:
        if not self._last_management_report:
            messagebox.showinfo("Management Report", "Generate a management report first.", parent=self)
            return
        path = filedialog.asksaveasfilename(
            title="Save Management Report", parent=self, defaultextension=".docx",
            filetypes=[("Word Document", "*.docx")],
        )
        if not path:
            return
        try:
            from docx import Document
            document = Document()
            document.add_heading("Campaign Management Report", level=1)
            profile = self.master_app.active_campaign_profile
            if profile:
                document.add_paragraph(f"Campaign: {profile.campaign_id} — {profile.campaign_name}")
            for key, value in self._last_management_report.items():
                document.add_heading(str(key).replace("_", " ").title(), level=2)
                if isinstance(value, list):
                    for item in value:
                        document.add_paragraph(str(item), style="List Bullet")
                elif isinstance(value, dict):
                    for sub_key, sub_value in value.items():
                        document.add_paragraph(f"{str(sub_key).replace('_', ' ').title()}: {sub_value}")
                else:
                    document.add_paragraph(str(value))
            document.save(path)
            messagebox.showinfo("Management Report", "Report saved successfully.", parent=self)
        except Exception as exc:
            messagebox.showerror("Management Report", str(exc), parent=self)

    # ---------------- Knowledge ----------------
    def _build_knowledge_tab(self) -> None:
        page = self._panel(self.knowledge_tab)
        card = ctk.CTkFrame(page)
        card.pack(fill="x", padx=8, pady=8)
        ctk.CTkLabel(card, text="Local Approved Knowledge", font=ctk.CTkFont(size=17, weight="bold")).pack(anchor="w", padx=12, pady=(12, 4))
        actions = ctk.CTkFrame(card, fg_color="transparent")
        actions.pack(fill="x", padx=12, pady=8)
        ctk.CTkButton(actions, text="Refresh", command=self._refresh_knowledge).pack(side="left", padx=(0, 6))
        ctk.CTkButton(actions, text="Selected Record AI Evidence", command=self._show_selected_evidence).pack(side="left", padx=6)
        ctk.CTkButton(actions, text="Clear AI Response Cache", command=self._clear_ai_cache).pack(side="left", padx=6)
        self.knowledge_output = self._output_box(card, 500)
        self.after(100, self._refresh_knowledge)

    def _refresh_knowledge(self) -> None:
        items = self.ai.knowledge.list(limit=300)
        self._set_box(self.knowledge_output, items)

    def _show_selected_evidence(self) -> None:
        row_id, _record = self._selected_row()
        if row_id is None:
            messagebox.showinfo("AI Evidence", "Select a lead row first.", parent=self)
            return
        evidence = self.ai.evidence.for_record(str(row_id + 1))
        self._set_box(self.knowledge_output, evidence or {"message": "No AI-derived evidence has been stored for this row."})

    def _clear_ai_cache(self) -> None:
        if not messagebox.askyesno("AI Cache", "Clear cached Qwen responses? Approved Local Knowledge will be kept.", parent=self):
            return
        count = self.ai.clear_cache()
        messagebox.showinfo("AI Cache", f"Cleared {count} cached response(s).", parent=self)
