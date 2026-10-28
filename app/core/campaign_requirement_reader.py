from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from app.core.campaign_models import CampaignRule, RuleSource
from app.core.document_reader import DocumentContent, DocumentReader


_LABEL_ALIASES: dict[str, tuple[str, str, str]] = {
    "country": ("Targeting", "allowed_countries", "in"),
    "countries": ("Targeting", "allowed_countries", "in"),
    "geo": ("Targeting", "allowed_countries", "in"),
    "geography": ("Targeting", "allowed_countries", "in"),
    "location": ("Targeting", "allowed_countries", "in"),
    "industry": ("Company", "allowed_industries", "in"),
    "industries": ("Company", "allowed_industries", "in"),
    "contact level": ("Prospect", "minimum_seniority", ">="),
    "seniority": ("Prospect", "minimum_seniority", ">="),
    "job level": ("Prospect", "minimum_seniority", ">="),
    "job titles": ("Prospect", "job_title_keywords", "contains_any"),
    "job title": ("Prospect", "job_title_keywords", "contains_any"),
    "title keywords": ("Prospect", "job_title_keywords", "contains_any"),
    "keywords": ("Prospect", "job_title_keywords", "contains_any"),
    "job function": ("Prospect", "job_functions", "in"),
    "job functions": ("Prospect", "job_functions", "in"),
    "department": ("Prospect", "job_functions", "in"),
    "departments": ("Prospect", "job_functions", "in"),
    "company size": ("Company", "company_size", "in"),
    "employee size": ("Company", "company_size", "in"),
    "employees": ("Company", "company_size", "in"),
    "revenue": ("Company", "revenue", "in"),
    "required fields": ("Validation", "required_fields", "contains_all"),
    "mandatory fields": ("Validation", "required_fields", "contains_all"),
    "lead goal": ("Campaign", "lead_goal", "equals"),
    "lead target": ("Campaign", "lead_goal", "equals"),
    "campaign id": ("Campaign", "campaign_id", "equals"),
    "campaign name": ("Campaign", "campaign_name", "equals"),
    "client": ("Campaign", "client", "equals"),
    "cpc": ("Campaign", "cpc", "equals"),
    "start date": ("Campaign", "start_date", "equals"),
    "end date": ("Campaign", "end_date", "equals"),
}

_LIST_FIELDS = {
    "allowed_countries",
    "allowed_industries",
    "job_title_keywords",
    "job_functions",
    "company_size",
    "revenue",
    "required_fields",
}


def _clean_label(value: object) -> str:
    text = str(value or "").strip().strip(":").strip()
    text = re.sub(r"\s+", " ", text)
    return text.casefold()


def _clean_value(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return re.sub(r"\s+", " ", text)


def _split_values(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []
    parts = re.split(r"\s*(?:,|;|\||\n|•|\u2022)\s*", text)
    return [part.strip(" -\t") for part in parts if part.strip(" -\t")]


def _is_label(text: str) -> tuple[str, str, str] | None:
    key = _clean_label(text)
    if key in _LABEL_ALIASES:
        return _LABEL_ALIASES[key]
    # tolerate labels with explanatory suffixes, e.g. "Job Titles / Keywords"
    for alias, mapping in _LABEL_ALIASES.items():
        if key.startswith(alias + " ") or key.startswith(alias + "/"):
            return mapping
    return None


@dataclass(slots=True)
class RequirementReadResult:
    rules: list[CampaignRule]
    warnings: list[str]


class CampaignRequirementReader:
    """Extract structured requirements from spreadsheet and Word specification files."""

    def __init__(self) -> None:
        self.document_reader = DocumentReader()

    def read(self, path: str | Path) -> RequirementReadResult:
        source = Path(path)
        suffix = source.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            return self._read_excel(source)
        if suffix == ".xls":
            return self._read_xls(source)
        if suffix in {".docx", ".doc"}:
            try:
                content = self.document_reader.read(source)
            except Exception as exc:
                return RequirementReadResult([], [f"{source.name}: {exc}"])
            return self._read_document_content(source, content)
        if suffix in {".csv", ".txt", ".tsv"}:
            return self._read_delimited_spec(source)
        return RequirementReadResult([], [f"Unsupported specification format: {source.suffix}"])

    def _read_excel(self, path: Path) -> RequirementReadResult:
        workbook = load_workbook(path, data_only=True, read_only=True)
        rules: list[CampaignRule] = []
        warnings: list[str] = []
        for worksheet in workbook.worksheets:
            # Build each column as an independent requirements stream. This handles
            # sheets like Specs.xlsx where Country and Industry occupy separate columns.
            max_rows = worksheet.max_row or 0
            max_cols = worksheet.max_column or 0
            for col in range(1, max_cols + 1):
                active_mapping: tuple[str, str, str] | None = None
                active_values: list[str] = []
                active_sources: list[RuleSource] = []

                def flush() -> None:
                    nonlocal active_mapping, active_values, active_sources
                    if not active_mapping or not active_values:
                        active_mapping = None
                        active_values = []
                        active_sources = []
                        return
                    category, field_name, operator = active_mapping
                    values = self._normalize_rule_values(field_name, active_values)
                    status, notes = self._default_rule_status(field_name, values)
                    value: Any = values if field_name in _LIST_FIELDS else (values[0] if values else "")
                    rules.append(
                        CampaignRule(
                            category=category,
                            field=field_name,
                            operator=operator,
                            value=value,
                            confidence=1.0,
                            status=status,
                            notes=notes,
                            sources=list(active_sources),
                        )
                    )
                    active_mapping = None
                    active_values = []
                    active_sources = []

                for row in range(1, max_rows + 1):
                    cell = worksheet.cell(row=row, column=col)
                    raw = _clean_value(cell.value)
                    if not raw:
                        continue
                    mapping = _is_label(raw)
                    if mapping:
                        flush()
                        active_mapping = mapping
                        # If a label contains an inline value after ':' parse it.
                        original = str(cell.value or "")
                        if ":" in original:
                            tail = original.split(":", 1)[1].strip()
                            if tail:
                                active_values.extend(_split_values(tail) or [tail])
                                active_sources.append(
                                    RuleSource(path.name, f"{worksheet.title}!{cell.coordinate}", original)
                                )
                        continue
                    if active_mapping:
                        active_values.append(raw)
                        active_sources.append(
                            RuleSource(path.name, f"{worksheet.title}!{cell.coordinate}", raw)
                        )
                flush()
        return RequirementReadResult(self._merge_compatible_rules(rules), warnings)

    def _read_xls(self, path: Path) -> RequirementReadResult:
        try:
            import xlrd
        except ImportError as exc:
            return RequirementReadResult([], [f"{path.name}: legacy XLS reading requires xlrd"] )
        book = xlrd.open_workbook(str(path), on_demand=True)
        rules: list[CampaignRule] = []
        for sheet in book.sheets():
            for col in range(sheet.ncols):
                active_mapping = None
                active_values: list[str] = []
                active_sources: list[RuleSource] = []
                def flush():
                    nonlocal active_mapping, active_values, active_sources
                    if not active_mapping or not active_values:
                        active_mapping = None; active_values = []; active_sources = []; return
                    category, field_name, operator = active_mapping
                    vals = self._normalize_rule_values(field_name, active_values)
                    status, notes = self._default_rule_status(field_name, vals)
                    rules.append(CampaignRule(category, field_name, operator, vals if field_name in _LIST_FIELDS else vals[0], confidence=0.98, status=status, notes=notes, sources=list(active_sources)))
                    active_mapping = None; active_values = []; active_sources = []
                for row in range(sheet.nrows):
                    raw = _clean_value(sheet.cell_value(row, col))
                    if not raw:
                        continue
                    mapping = _is_label(raw)
                    if mapping:
                        flush(); active_mapping = mapping; continue
                    if active_mapping:
                        active_values.append(raw)
                        active_sources.append(RuleSource(path.name, f"{sheet.name}!R{row+1}C{col+1}", raw))
                flush()
        return RequirementReadResult(self._merge_compatible_rules(rules), [])

    def _read_document_content(self, path: Path, content: DocumentContent) -> RequirementReadResult:
        rules: list[CampaignRule] = []
        warnings: list[str] = []
        lines: list[tuple[str, str]] = []
        for index, line in enumerate(content.lines, start=1):
            lines.append((line, f"Paragraph {index}"))
        for table_index, table in enumerate(content.tables, start=1):
            for row_index, row in enumerate(table, start=1):
                cells = [cell.strip() for cell in row if cell.strip()]
                if not cells:
                    continue
                # Treat 2-column tables as label/value pairs first.
                if len(cells) >= 2 and _is_label(cells[0]):
                    mapping = _is_label(cells[0])
                    assert mapping is not None
                    category, field_name, operator = mapping
                    vals = []
                    for cell in cells[1:]:
                        vals.extend(_split_values(cell) or [cell])
                    vals = self._normalize_rule_values(field_name, vals)
                    status, notes = self._default_rule_status(field_name, vals)
                    rules.append(
                        CampaignRule(
                            category, field_name, operator,
                            vals if field_name in _LIST_FIELDS else (vals[0] if vals else ""),
                            confidence=0.96,
                            status=status,
                            notes=notes,
                            sources=[RuleSource(path.name, f"Table {table_index}, row {row_index}", " | ".join(cells))],
                        )
                    )
                else:
                    lines.append((" | ".join(cells), f"Table {table_index}, row {row_index}"))

        active_mapping: tuple[str, str, str] | None = None
        active_values: list[str] = []
        active_sources: list[RuleSource] = []

        def flush() -> None:
            nonlocal active_mapping, active_values, active_sources
            if not active_mapping or not active_values:
                active_mapping = None
                active_values = []
                active_sources = []
                return
            category, field_name, operator = active_mapping
            vals = self._normalize_rule_values(field_name, active_values)
            status, notes = self._default_rule_status(field_name, vals)
            rules.append(
                CampaignRule(
                    category, field_name, operator,
                    vals if field_name in _LIST_FIELDS else (vals[0] if vals else ""),
                    confidence=0.94,
                    status=status,
                    notes=notes,
                    sources=list(active_sources),
                )
            )
            active_mapping = None
            active_values = []
            active_sources = []

        for raw, location in lines:
            text = raw.strip().strip("•\t")
            if not text:
                continue
            # Inline "Label: value" form.
            if ":" in text:
                left, right = text.split(":", 1)
                mapping = _is_label(left)
                if mapping:
                    flush()
                    if right.strip():
                        category, field_name, operator = mapping
                        vals = self._normalize_rule_values(field_name, _split_values(right) or [right.strip()])
                        status, notes = self._default_rule_status(field_name, vals)
                        rules.append(
                            CampaignRule(
                                category, field_name, operator,
                                vals if field_name in _LIST_FIELDS else (vals[0] if vals else ""),
                                confidence=0.98,
                                status=status,
                                notes=notes,
                                sources=[RuleSource(path.name, location, text)],
                            )
                        )
                    else:
                        active_mapping = mapping
                    continue
            mapping = _is_label(text)
            if mapping:
                flush()
                active_mapping = mapping
                continue
            if active_mapping:
                # Stop a list when another obvious section heading appears.
                if len(text) < 60 and text.endswith(":"):
                    flush()
                    continue
                active_values.extend(_split_values(text) or [text])
                active_sources.append(RuleSource(path.name, location, text))
        flush()
        return RequirementReadResult(self._merge_compatible_rules(rules), warnings)

    def _read_delimited_spec(self, path: Path) -> RequirementReadResult:
        # Small textual spec CSVs are handled as rows of text/label-value pairs.
        warnings: list[str] = []
        rows: list[list[str]] = []
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                with path.open("r", encoding=encoding, newline="") as handle:
                    sample = handle.read(4096)
                    handle.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                        reader = csv.reader(handle, dialect)
                    except csv.Error:
                        reader = csv.reader(handle, delimiter=",")
                    rows = [[cell.strip() for cell in row] for row in reader]
                break
            except UnicodeDecodeError:
                continue
        if not rows:
            return RequirementReadResult([], [f"Could not read {path.name}"])
        # Convert rows to a temporary document-like representation.
        lines = []
        tables = [rows]
        content = DocumentContent(str(path), lines=lines, tables=tables, backend="delimited")
        return self._read_document_content(path, content)

    @staticmethod
    def _normalize_rule_values(field_name: str, values: Iterable[str]) -> list[str]:
        cleaned: list[str] = []
        seen: set[str] = set()
        for value in values:
            item = _clean_value(value).strip(" :-")
            if not item:
                continue
            # Preserve source wording while deduplicating case-insensitively.
            key = item.casefold()
            if key not in seen:
                seen.add(key)
                cleaned.append(item)
        if field_name == "minimum_seniority" and cleaned:
            return [cleaned[0]]
        return cleaned

    @staticmethod
    def _default_rule_status(field_name: str, values: list[str]) -> tuple[str, str]:
        # CPC is captured faithfully but its validation meaning varies by client.
        if field_name == "cpc":
            return "Needs Review", "Extracted from source; validation behavior is not assumed automatically."
        if not values:
            return "Needs Review", "No usable value was extracted."
        return "Extracted", ""

    @staticmethod
    def _merge_compatible_rules(rules: list[CampaignRule]) -> list[CampaignRule]:
        merged: dict[str, CampaignRule] = {}
        output: list[CampaignRule] = []
        for rule in rules:
            if rule.field not in _LIST_FIELDS:
                output.append(rule)
                continue
            key = rule.key
            current = merged.get(key)
            if current is None:
                current = CampaignRule(
                    rule.category, rule.field, rule.operator, [],
                    confidence=rule.confidence,
                    status=rule.status,
                    notes=rule.notes,
                    sources=[],
                )
                merged[key] = current
                output.append(current)
            existing = {str(v).casefold() for v in current.value}
            for value in rule.value:
                if str(value).casefold() not in existing:
                    current.value.append(value)
                    existing.add(str(value).casefold())
            current.sources.extend(rule.sources)
            current.confidence = min(current.confidence, rule.confidence)
        return output
