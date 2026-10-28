from __future__ import annotations

import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path


@dataclass(slots=True)
class DocumentContent:
    path: str
    lines: list[str] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)
    backend: str = ""

    @property
    def text(self) -> str:
        parts = list(self.lines)
        for table in self.tables:
            for row in table:
                parts.append(" | ".join(cell for cell in row if cell))
        return "\n".join(part for part in parts if part.strip())


class DocumentReader:
    """Read campaign requirements from Word documents.

    DOCX uses python-docx. Legacy DOC files use Microsoft Word automation on
    Windows when available, with antiword as a command-line fallback. PDF is
    intentionally unsupported for the campaign workflow.
    """

    SUPPORTED = {".docx", ".doc"}

    def read(self, path: str | Path) -> DocumentContent:
        source = Path(path)
        suffix = source.suffix.lower()
        if suffix == ".docx":
            return self._read_docx(source)
        if suffix == ".doc":
            return self._read_doc(source)
        raise ValueError(f"Unsupported document type for Campaign Manager: {suffix}")

    def _read_docx(self, path: Path) -> DocumentContent:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX reading requires python-docx. Install requirements.txt.") from exc

        document = Document(str(path))
        lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        tables: list[list[list[str]]] = []
        for table in document.tables:
            rows: list[list[str]] = []
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            if rows:
                tables.append(rows)
        return DocumentContent(str(path), lines=lines, tables=tables, backend="python-docx")

    def _read_doc(self, path: Path) -> DocumentContent:
        # Best Windows path: Word itself can accurately read legacy binary .doc.
        try:
            import win32com.client  # type: ignore

            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            doc = None
            try:
                doc = word.Documents.Open(str(path.resolve()), ReadOnly=True)
                raw = str(doc.Content.Text or "")
            finally:
                if doc is not None:
                    doc.Close(False)
                word.Quit()
            lines = [line.strip() for line in raw.replace("\r", "\n").split("\n") if line.strip()]
            return DocumentContent(str(path), lines=lines, backend="Microsoft Word COM")
        except Exception:
            pass

        antiword = shutil.which("antiword")
        if antiword:
            result = subprocess.run(
                [antiword, str(path)], capture_output=True, text=True, timeout=60, check=False
            )
            if result.returncode == 0 and result.stdout.strip():
                lines = [line.strip() for line in result.stdout.splitlines() if line.strip()]
                return DocumentContent(str(path), lines=lines, backend="antiword")

        raise RuntimeError(
            "Legacy .doc file detected, but no reader is available. On Windows, install Microsoft Word "
            "and pywin32 (included in this project's Windows requirements), or install antiword."
        )
