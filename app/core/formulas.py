"""Local, single-worksheet formulas. User expressions are never Python code.

The worksheet keeps expressions separately from the values used by the rest of
the CRM. Excel compatibility is bounded by the installed formulas engine.
"""
from __future__ import annotations

from functools import lru_cache
from graphlib import TopologicalSorter, CycleError
import importlib
import math
import re
import functools

import numpy as np
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils import get_column_letter, column_index_from_string

from app.constants import MISSING_VALUE

MAX_RANGE_CELLS = 100_000
PROTECTED_FIELDS = {"QA Status", "Primary Reason"}
STRINGS = re.compile(r'("(?:[^"]|"")*")')
CELL = re.compile(r"^(\$?[A-Za-z]{1,3})(\$?[1-9][0-9]*)$")
COLUMN_RANGE = re.compile(r"^\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})$")
ROW_RANGE = re.compile(r"^\$?([1-9][0-9]*):\$?([1-9][0-9]*)$")
NUMERIC = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")


class FormulaError(ValueError):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


@lru_cache(maxsize=1)
def function_catalog():
    import formulas
    import formulas.functions as functions
    labels = {"info": "Information", "logic": "Logical", "math": "Math & trigonometry",
              "stat": "Statistical", "financial": "Financial", "text": "Text",
              "look": "Lookup & reference", "eng": "Engineering", "date": "Date & time",
              "comp": "Compatibility"}
    result = {}
    supported = formulas.get_functions()
    for module, category in labels.items():
        for name in importlib.import_module("formulas.functions." + module).FUNCTIONS:
            if isinstance(name, str) and not name.startswith("_") and name in supported:
                result.setdefault(name, category)
    # These are parsed language forms rather than ordinary registry callables.
    result.update({"LET": "Logical", "LAMBDA": "Logical"})
    return dict(sorted(result.items()))


def _outside_strings(expression, transform):
    return "".join(part if i % 2 else transform(part) for i, part in enumerate(STRINGS.split(expression)))


def canonical_formula(expression, row, column, headers, *, resolve_context=True):
    """Resolve [Header], simple header names, and columnN to source-row A1 refs."""
    if len(expression) > 8192:
        raise FormulaError("#VALUE!", "A formula can contain at most 8,192 characters.")
    by_name = {name.casefold(): i for i, name in enumerate(headers)}

    def transform(part):
        if "!" in part.replace("#REF!", "").replace("#VALUE!", "").replace("#DIV/0!", "").replace("#NUM!", "").replace("#NULL!", ""):
            raise FormulaError("#REF!", "External workbooks and other sheets are not supported. Use this worksheet's columns.")
        def header(match):
            name = match.group(1).strip().casefold()
            if name not in by_name:
                raise FormulaError("#REF!", f"Unknown column: {match.group(1)}")
            return f"{get_column_letter(by_name[name] + 1)}{row + 1}"
        part = re.sub(r"\[([^\[\]]+)\]", header, part)
        if "[" in part or "]" in part or "|" in part:
            raise FormulaError("#REF!", "Use [Column name] for a worksheet column; external links are not supported.")
        def word(match):
            name = match.group(0)
            tail = part[match.end():].lstrip()
            if tail.startswith("(") or name.upper() in {"TRUE", "FALSE"}:
                return name.upper()
            if name.casefold() in by_name and not CELL.fullmatch(name):
                return f"{get_column_letter(by_name[name.casefold()] + 1)}{row + 1}"
            positional = re.fullmatch(r"column([1-9]\d*)", name, re.I)
            if positional:
                number = int(positional[1])
                if number > len(headers):
                    raise FormulaError("#REF!", f"Column {number} does not exist.")
                return f"{get_column_letter(number)}{row + 1}"
            return name
        part = re.sub(r"(?<![\w.$#])[A-Za-z_][A-Za-z_0-9.]*(?![\w])", word, part)
        if resolve_context:
            part = re.sub(r"\bROW\(\s*\)", str(row + 1), part, flags=re.I)
            part = re.sub(r"\bCOLUMN\(\s*\)", str(column + 1), part, flags=re.I)
        return part
    return _outside_strings(expression, transform)


def reference_cells(reference, headers, record_count):
    """Return rectangular coordinates, bounded before allocating any arrays."""
    whole = COLUMN_RANGE.fullmatch(reference)
    rows = ROW_RANGE.fullmatch(reference)
    if whole:
        c1, c2 = (column_index_from_string(v.upper()) - 1 for v in whole.groups())
        r1, r2 = 0, max(0, record_count - 1)
    elif rows:
        r1, r2 = (int(v) - 1 for v in rows.groups())
        c1, c2 = 0, len(headers) - 1
    else:
        parts = reference.split(":")
        if len(parts) > 2 or not all(CELL.fullmatch(p) for p in parts):
            raise FormulaError("#NAME?", f"Unknown reference: {reference}. Use [Header], column1, or A1.")
        first, last = CELL.fullmatch(parts[0]), CELL.fullmatch(parts[-1])
        c1, c2 = (column_index_from_string(m[1].replace("$", "").upper()) - 1 for m in (first, last))
        r1, r2 = (int(m[2].replace("$", "")) - 1 for m in (first, last))
    if min(r1, c1) < 0 or r2 < r1 or c2 < c1 or c2 >= len(headers) or r2 >= 1048576:
        raise FormulaError("#REF!", f"Reference is outside the worksheet: {reference}")
    if (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_RANGE_CELLS:
        raise FormulaError("#NUM!", f"A range can contain at most {MAX_RANGE_CELLS:,} cells. Use a smaller range.")
    if any(h in PROTECTED_FIELDS for h in headers[c1:c2 + 1]):
        raise FormulaError("#REF!", "QA Status and Primary Reason are validation outputs and cannot be formula inputs.")
    return r1, r2, c1, c2


@lru_cache(maxsize=1024)
def compile_formula(expression):
    import formulas
    from formulas.tokens.function import Function
    _install_limits()
    # Nested LET/LAMBDA parsing may compile children immediately. Reject unknown
    # calls before parsing so unsupported names never reach the engine dispatcher.
    for name in re.findall(r"\b([A-Za-z_][\w.]*)\s*\(", STRINGS.sub('""', expression)):
        if name.upper().removeprefix("_XLFN.") not in function_catalog():
            raise FormulaError("#NAME?", f"{name.upper()} is not supported. Open Insert function to see available functions.")
    tokens, builder = formulas.Parser().ast(expression)
    for token in tokens:
        if isinstance(token, Function):
            name = token.attr["name"].upper()
            if name.removeprefix("_XLFN.") not in function_catalog() and name not in {"ARRAY", "ARRAYROW"}:
                raise FormulaError("#NAME?", f"{name} is not supported. Open Insert function to see available functions.")
    return builder.compile()


@lru_cache(maxsize=1)
def _install_limits():
    """Bound functions which can allocate huge results from a single scalar."""
    import formulas
    from formulas.functions import Error
    registry = formulas.get_functions()
    def guard(name, check):
        original = registry[name]
        @functools.wraps(original)
        def limited(*args, **kwargs):
            try:
                arrays = [np.asarray(getattr(a, "value", a), dtype=object) for a in args]
                if not check(arrays):
                    return Error.errors["#NUM!"]
            except (TypeError, ValueError, IndexError, OverflowError):
                return Error.errors["#VALUE!"]
            return original(*args, **kwargs)
        registry[name] = limited
        if "_XLFN." + name in registry:
            registry["_XLFN." + name] = limited
    guard("REPT", lambda a: all(0 <= int(float(n)) and len(str(t)) * int(float(n)) <= 32767
                               for t, n in zip(*(v.flat for v in np.broadcast_arrays(*a)))))
    guard("MUNIT", lambda a: all(0 < float(n) ** 2 <= MAX_RANGE_CELLS for n in a[0].flat))
    # The upstream NumPy ufunc wrapper cannot return a matrix with NumPy 2.x.
    # A scalar dimension produces the same bounded identity matrix directly.
    def unit_matrix(value):
        try:
            dimension = int(np.asarray(getattr(value, "value", value)).item())
            if dimension < 1 or dimension ** 2 > MAX_RANGE_CELLS:
                return Error.errors["#NUM!"]
            return np.identity(dimension, dtype=object)
        except (ValueError, TypeError, OverflowError):
            return Error.errors["#VALUE!"]
    registry["MUNIT"] = registry["_XLFN.MUNIT"] = unit_matrix
    guard("MAKEARRAY", lambda a: float(a[0].item()) * float(a[1].item()) <= MAX_RANGE_CELLS)
    guard("MMULT", lambda a: a[0].shape[0] * a[1].shape[-1] <= MAX_RANGE_CELLS)


def cell_value(value):
    from formulas.functions import Error
    import schedula
    if value is None or str(value) in {"", MISSING_VALUE}:
        return schedula.EMPTY
    if str(value) in Error.errors:
        return Error.errors[str(value)]
    if str(value) in {"#CYCLE!", "#SPILL!"}:
        return Error.errors["#REF!"]
    if str(value).upper() in {"TRUE", "FALSE"}:
        return str(value).upper() == "TRUE"
    if NUMERIC.fullmatch(str(value).strip()):
        number = float(value)
        if math.isfinite(number):
            return number
    text = str(value)
    if len(text) > 32767:
        raise FormulaError("#VALUE!", "A calculated text value can contain at most 32,767 characters.")
    return text


def display_result(value):
    import schedula
    array = np.asarray(value, dtype=object)
    if array.size != 1:
        raise FormulaError("#SPILL!", "This formula returns multiple cells. Wrap it in INDEX, SUM or TEXTJOIN to return one value.")
    value = array.item()
    if value is schedula.EMPTY or value is None:
        return "0"
    if isinstance(value, (bool, np.bool_)):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float, np.number)):
        if not math.isfinite(float(value)):
            raise FormulaError("#NUM!", "The result is not a finite number.")
        return format(float(value), ".15g")
    text = str(value)
    if len(text) > 32767:
        raise FormulaError("#VALUE!", "A calculated text value can contain at most 32,767 characters.")
    return text


def recalculate(manager):
    """Evaluate in dependency order, returning cells whose calculated value changed."""
    from formulas.ranges import Ranges
    headers = manager.headers
    formula_cells = set(manager.formulas)
    plans, graph, values, errors = {}, {}, {}, {}
    for cell, expression in manager.formulas.items():
        row, header = cell
        try:
            formula = canonical_formula(expression, row, headers.index(header), headers)
            function = compile_formula(formula)
            refs = {ref: reference_cells(str(ref), headers, len(manager.records)) for ref in function.inputs}
            dependencies = set()
            for r1, r2, c1, c2 in refs.values():
                dependencies.update((r, headers[c]) for r in range(r1, min(r2 + 1, len(manager.records)))
                                    for c in range(c1, c2 + 1) if (r, headers[c]) in formula_cells)
            graph[cell] = dependencies
            plans[cell] = (function, refs)
        except Exception as exc:
            values[cell] = exc.code if isinstance(exc, FormulaError) else "#VALUE!"
            errors[cell] = str(exc) if isinstance(exc, FormulaError) else "Check the formula syntax and arguments."
            graph[cell] = set()
    order = TopologicalSorter(graph)
    try:
        order.prepare()
    except CycleError:
        pass
    while order.is_active():
        ready = order.get_ready()
        for cell in ready:
            if cell not in values:
                function, refs = plans[cell]
                try:
                    arguments = []
                    for r1, r2, c1, c2 in refs.values():
                        array = [[cell_value(values.get((r, headers[c]), manager.records[r].get(headers[c], MISSING_VALUE))
                                             if r < len(manager.records) else MISSING_VALUE)
                                  for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]
                        bounded_ref = f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2 + 1)}{r2 + 1}"
                        arguments.append(Ranges().push(bounded_ref, np.array(array, dtype=object)))
                    with np.errstate(all="ignore"):
                        values[cell] = display_result(function(*arguments))
                    if values[cell].startswith("#"):
                        errors[cell] = "Excel calculation error. Check the inputs and function arguments."
                except Exception as exc:
                    values[cell] = exc.code if isinstance(exc, FormulaError) else "#VALUE!"
                    errors[cell] = str(exc) if isinstance(exc, FormulaError) else "The function could not calculate with these arguments."
            order.done(cell)
    for cell in formula_cells - values.keys():
        values[cell] = "#CYCLE!"
        errors[cell] = "Circular reference: this formula depends on itself or another circular formula."
    changed = set()
    for (row, header), value in values.items():
        if manager.records[row].get(header) != value:
            manager.records[row][header] = value
            changed.add((row, header))
    manager.formula_errors = errors
    return changed


def translate_formula(expression, source, target, headers):
    formula = canonical_formula(expression, source[0], source[1], headers, resolve_context=False)
    origin = f"{get_column_letter(source[1] + 1)}{source[0] + 1}"
    destination = f"{get_column_letter(target[1] + 1)}{target[0] + 1}"
    try:
        return Translator(formula, origin=origin).translate_formula(destination)
    except TranslatorError:
        return "=#REF!"


def remove_column_references(expression, row, column, headers, deleted):
    """Keep positional references attached to surviving columns after deletion."""
    try:
        formula = canonical_formula(expression, row, column, headers, resolve_context=False)
        tokens = Tokenizer(formula).items
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            r1, r2, c1, c2 = reference_cells(token.value, headers, 1)
            surviving = [c for c in range(c1, c2 + 1) if c not in deleted]
            if not surviving:
                token.value = "#REF!"
                continue
            parts = token.value.split(":")
            for i, old in enumerate((surviving[0], surviving[-1])[:len(parts)]):
                new = old - sum(c < old for c in deleted)
                parts[i] = re.sub(r"[A-Za-z]+", get_column_letter(new + 1), parts[i])
            token.value = ":".join(parts)
        return "=" + "".join(token.value for token in tokens)
    except Exception:
        return expression


def remove_row_references(expression, row, column, headers, deleted):
    try:
        tokens = Tokenizer(canonical_formula(expression, row, column, headers, resolve_context=False)).items
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE" or COLUMN_RANGE.fullmatch(token.value):
                continue
            r1, r2, c1, c2 = reference_cells(token.value, headers, 1)
            surviving = [r for r in range(r1, r2 + 1) if r not in deleted]
            if not surviving:
                token.value = "#REF!"
                continue
            parts = token.value.split(":")
            for i, old in enumerate((surviving[0], surviving[-1])[:len(parts)]):
                parts[i] = re.sub(r"\d+", str(old - sum(r < old for r in deleted) + 1), parts[i])
            token.value = ":".join(parts)
        return "=" + "".join(token.value for token in tokens)
    except Exception:
        return expression
