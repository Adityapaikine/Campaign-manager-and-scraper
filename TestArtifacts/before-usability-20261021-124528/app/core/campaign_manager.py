from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import sqlite3
import zipfile
from collections import Counter
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from app.core.cache_db import initialize_database
from app.core.campaign_models import (
    CampaignConflict,
    CampaignFileInfo,
    CampaignProfile,
    CampaignRule,
    RuleSource,
)
from app.core.campaign_requirement_reader import CampaignRequirementReader
from app.utils.config import CAMPAIGN_DIR, DB_PATH, ensure_directories

SUPPORTED_CAMPAIGN_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".txt", ".docx", ".doc"}


def _safe_name(text: str) -> str:
    text = re.sub(r"[^A-Za-z0-9._ -]+", "_", text).strip(" ._")
    return text[:120] or "campaign"


def _normalize_header(text: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(text or "").casefold())


def _decode_sample(path: Path, size: int = 8192) -> str:
    data = path.read_bytes()[:size]
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _csv_header_and_count(path: Path) -> tuple[list[str], int, bool]:
    text = _decode_sample(path)
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return [], 0, False
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:20]), delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    first = next(csv.reader([lines[0]], delimiter=delimiter), [])
    header = [cell.strip().strip('"') for cell in first]
    header_norm = {_normalize_header(cell) for cell in header}
    known_header_tokens = {
        "email", "emailaddress", "businessemailid", "domain", "domainname", "company",
        "companyname", "website", "webaddress", "firstname", "lastname", "title",
        "country", "industry", "domainlimit", "linkedinleadvalidationurl",
    }
    looks_headerless = not bool(header_norm & known_header_tokens)
    if len(header) == 1:
        value = header[0].strip()
        if "@" in value or ("." in value and " " not in value and len(value) < 150):
            looks_headerless = True
    # Fast line count; quoted multiline fields are uncommon in reference lists. The
    # count is metadata only and not used to parse/validate records.
    with path.open("rb") as handle:
        count = sum(chunk.count(b"\n") for chunk in iter(lambda: handle.read(1024 * 1024), b""))
    record_count = max(0, count - (0 if looks_headerless else 1))
    return header, record_count, looks_headerless


def _xlsx_header_and_count(path: Path) -> tuple[list[str], int]:
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        sheet = book.sheet_by_index(0)
        columns = [str(sheet.cell_value(0, col)).strip() for col in range(sheet.ncols)] if sheet.nrows else []
        return [c for c in columns if c], max(0, sheet.nrows - 1)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 5), values_only=True)
    first = next(rows, ())
    columns = [str(v).strip() for v in first if v is not None and str(v).strip()]
    return columns, max(0, (ws.max_row or 0) - 1)


class CampaignFileClassifier:
    def classify(self, path: Path) -> CampaignFileInfo:
        suffix = path.suffix.lower()
        name = path.name.casefold()
        columns: list[str] = []
        count = 0
        headerless = False
        notes: list[str] = []
        if suffix in {".csv", ".tsv", ".txt"}:
            try:
                columns, count, headerless = _csv_header_and_count(path)
                if headerless:
                    notes.append("Headerless/single-column source detected; Campaign Manager will infer the field type.")
            except Exception as exc:
                notes.append(f"Could not inspect delimited file: {exc}")
        elif suffix in {".xlsx", ".xlsm", ".xls"}:
            try:
                columns, count = _xlsx_header_and_count(path)
            except Exception as exc:
                notes.append(f"Could not inspect workbook: {exc}")

        norm_cols = {_normalize_header(c) for c in columns}
        role = "unknown"
        confidence = 0.55

        if suffix in {".doc", ".docx"}:
            role, confidence = "specification", 0.90
        elif any(token in name for token in ("spec", "brief", "instruction", "requirement", "campaign detail")):
            role, confidence = "specification", 0.97
        elif "competitor" in name:
            role, confidence = "competitor_exclusion", 0.98
        elif "domain" in name and "supp" in name:
            if "domainlimit" in norm_cols or any("limit" in c for c in norm_cols):
                role, confidence = "domain_limit", 0.99
            else:
                role, confidence = "domain_suppression", 0.96
        elif "processed" in name and ("supp" in name or "lead" in name):
            role, confidence = "processed_suppression", 0.98
        elif "internal" in name and "supp" in name:
            role, confidence = "internal_suppression", 0.99
        elif "reference" in name and "tal" in name:
            role, confidence = "reference_tal", 0.99
        elif "tal" in name:
            role, confidence = "tal", 0.98
        elif "supp" in name:
            role, confidence = "client_suppression", 0.95
        elif "domainlimit" in norm_cols:
            role, confidence = "domain_limit", 0.95
        elif norm_cols and norm_cols <= {"domain", "domainname", "website", "webaddress", "companywebsite"}:
            role, confidence = "tal", 0.72
        elif norm_cols and norm_cols <= {"email", "emailaddress", "businessemail", "businessemailid", "workemail"}:
            role, confidence = "client_suppression", 0.72

        if headerless and role == "unknown" and suffix in {".csv", ".txt"}:
            sample = _decode_sample(path, 1024).strip().splitlines()
            first = sample[0].strip().strip('"') if sample else ""
            if "@" in first:
                role, confidence = "client_suppression", 0.80
            elif "." in first:
                role, confidence = "tal", 0.68

        return CampaignFileInfo(
            relative_path=path.name,
            absolute_path=str(path),
            role=role,
            format=suffix.lstrip("."),
            record_count=count,
            confidence=confidence,
            notes=" ".join(notes),
            columns=columns,
            metadata={"headerless": headerless},
        )


class CampaignProfileStore:
    def __init__(self) -> None:
        initialize_database()

    def save(self, profile: CampaignProfile) -> None:
        payload = json.dumps(profile.to_dict(), ensure_ascii=False)
        with sqlite3.connect(DB_PATH) as conn:
            conn.execute(
                """
                INSERT INTO campaign_profiles(campaign_id, campaign_name, source_root, payload_json, active, updated_at)
                VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(campaign_id) DO UPDATE SET
                    campaign_name=excluded.campaign_name,
                    source_root=excluded.source_root,
                    payload_json=excluded.payload_json,
                    active=excluded.active,
                    updated_at=CURRENT_TIMESTAMP
                """,
                (profile.campaign_id, profile.campaign_name, profile.source_root, payload, int(profile.active)),
            )
            if profile.active:
                conn.execute("UPDATE campaign_profiles SET active=0 WHERE campaign_id <> ?", (profile.campaign_id,))
            conn.commit()

    def list(self) -> list[CampaignProfile]:
        with sqlite3.connect(DB_PATH) as conn:
            rows = conn.execute(
                "SELECT payload_json, active FROM campaign_profiles ORDER BY updated_at DESC"
            ).fetchall()
        result: list[CampaignProfile] = []
        for payload, active in rows:
            try:
                profile = CampaignProfile.from_dict(json.loads(payload))
                profile.active = bool(active)
                result.append(profile)
            except Exception:
                continue
        return result

    def get(self, campaign_id: str) -> CampaignProfile | None:
        with sqlite3.connect(DB_PATH) as conn:
            row = conn.execute(
                "SELECT payload_json, active FROM campaign_profiles WHERE campaign_id=?", (campaign_id,)
            ).fetchone()
        if not row:
            return None
        profile = CampaignProfile.from_dict(json.loads(row[0]))
        profile.active = bool(row[1])
        return profile

    def active(self) -> CampaignProfile | None:
        with sqlite3.connect(DB_PATH) as conn:
            row = conn.execute(
                "SELECT payload_json FROM campaign_profiles WHERE active=1 ORDER BY updated_at DESC LIMIT 1"
            ).fetchone()
        return CampaignProfile.from_dict(json.loads(row[0])) if row else None

    def activate(self, campaign_id: str) -> CampaignProfile | None:
        with sqlite3.connect(DB_PATH) as conn:
            conn.execute("UPDATE campaign_profiles SET active=0")
            conn.execute("UPDATE campaign_profiles SET active=1, updated_at=CURRENT_TIMESTAMP WHERE campaign_id=?", (campaign_id,))
            conn.commit()
        profile = self.get(campaign_id)
        if profile:
            profile.active = True
            self.save(profile)
        return profile

    def delete(self, campaign_id: str) -> bool:
        """Delete a saved campaign profile and return whether it was active.

        This intentionally deletes only the application's saved profile. Original
        source folders/ZIPs and imported campaign package files are left untouched.
        """
        campaign_id = str(campaign_id or "").strip()
        if not campaign_id:
            return False
        with sqlite3.connect(DB_PATH) as conn:
            row = conn.execute(
                "SELECT active FROM campaign_profiles WHERE campaign_id=?",
                (campaign_id,),
            ).fetchone()
            if row is None:
                return False
            was_active = bool(row[0])
            conn.execute("DELETE FROM campaign_profiles WHERE campaign_id=?", (campaign_id,))
            conn.commit()
        return was_active


class CampaignPackageImporter:
    def __init__(self) -> None:
        self.classifier = CampaignFileClassifier()
        self.requirement_reader = CampaignRequirementReader()
        self.store = CampaignProfileStore()

    def import_package(self, source: str | Path) -> CampaignProfile:
        ensure_directories()
        src = Path(source).resolve()
        if not src.exists():
            raise FileNotFoundError(src)
        destination = self._copy_or_extract(src)
        campaign_id, campaign_name = self._campaign_identity(destination.name)
        files: list[CampaignFileInfo] = []
        rules: list[CampaignRule] = []
        warnings: list[str] = []

        for path in sorted(destination.rglob("*")):
            if not path.is_file():
                continue
            if path.suffix.lower() not in SUPPORTED_CAMPAIGN_EXTENSIONS:
                continue
            info = self.classifier.classify(path)
            info.relative_path = str(path.relative_to(destination))
            files.append(info)
            if info.role == "specification":
                result = self.requirement_reader.read(path)
                rules.extend(result.rules)
                warnings.extend(result.warnings)
            elif info.role == "domain_limit":
                try:
                    headers, rows = CampaignReferenceLoader.read_rows(info)
                    domain_col = next((h for h in headers if "domain" in _normalize_header(h)), headers[0] if headers else "")
                    limit_col = next((h for h in headers if "limit" in _normalize_header(h)), "")
                    limits: dict[str, int | float | str] = {}
                    if domain_col and limit_col:
                        for row in rows:
                            domain = str(row.get(domain_col) or "").strip().casefold()
                            raw_limit = str(row.get(limit_col) or "").strip()
                            if not domain or not raw_limit:
                                continue
                            try:
                                numeric = float(raw_limit)
                                limit_value = int(numeric) if numeric.is_integer() else numeric
                            except ValueError:
                                limit_value = raw_limit
                            limits[domain] = limit_value
                    if limits:
                        rules.append(
                            CampaignRule(
                                category="Company",
                                field="domain_limits",
                                operator="max_per_domain",
                                value=limits,
                                confidence=0.98,
                                status="Needs Review",
                                notes="Domain/limit pairs were found explicitly, but their operational counting semantics require user confirmation before enforcement.",
                                sources=[RuleSource(info.relative_path, "Domain Limit table", f"{len(limits)} domain limit entries")],
                            )
                        )
                except Exception as exc:
                    warnings.append(f"Could not read domain limits from {info.relative_path}: {exc}")

        rules = self._merge_rules(rules)
        conflicts = self._find_conflicts(rules)
        self._apply_identity_rules(rules, campaign_id, campaign_name)
        notes = "\n".join(warnings)
        profile = CampaignProfile(
            campaign_id=campaign_id,
            campaign_name=campaign_name,
            source_root=str(destination),
            files=files,
            rules=rules,
            conflicts=conflicts,
            notes=notes,
        )
        self.store.save(profile)
        return profile

    def _copy_or_extract(self, source: Path) -> Path:
        if source.is_file() and source.suffix.lower() == ".zip":
            base_name = _safe_name(source.stem)
            digest = hashlib.sha1(str(source).encode("utf-8")).hexdigest()[:8]
            destination = CAMPAIGN_DIR / f"{base_name}_{digest}"
            if destination.exists():
                shutil.rmtree(destination)
            destination.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(source, "r") as archive:
                root = destination.resolve()
                for member in archive.infolist():
                    target = (destination / member.filename).resolve()
                    if root not in target.parents and target != root:
                        raise ValueError("Unsafe path found inside campaign ZIP.")
                archive.extractall(destination)
            children = [p for p in destination.iterdir() if p.name != "__MACOSX"]
            if len(children) == 1 and children[0].is_dir():
                return children[0]
            return destination

        if source.is_dir():
            base_name = _safe_name(source.name)
            digest = hashlib.sha1(str(source).encode("utf-8")).hexdigest()[:8]
            destination = CAMPAIGN_DIR / f"{base_name}_{digest}"
            if destination.exists():
                shutil.rmtree(destination)
            shutil.copytree(source, destination)
            return destination
        raise ValueError("Campaign source must be a folder or ZIP archive.")

    @staticmethod
    def _campaign_identity(folder_name: str) -> tuple[str, str]:
        match = re.match(r"\s*(\d{3,})\s*[-_–—]\s*(.+?)\s*$", folder_name)
        if match:
            return match.group(1), match.group(2).strip()
        digits = re.search(r"\b(\d{3,})\b", folder_name)
        campaign_id = digits.group(1) if digits else _safe_name(folder_name).replace(" ", "_")
        return campaign_id, folder_name.strip()

    @staticmethod
    def _apply_identity_rules(rules: list[CampaignRule], campaign_id: str, campaign_name: str) -> None:
        fields = {rule.field for rule in rules}
        if "campaign_id" not in fields:
            rules.insert(0, CampaignRule("Campaign", "campaign_id", "equals", campaign_id, confidence=0.98, status="Confirmed", notes="Derived from campaign folder/package name."))
        if "campaign_name" not in fields:
            rules.insert(1, CampaignRule("Campaign", "campaign_name", "equals", campaign_name, confidence=0.95, status="Confirmed", notes="Derived from campaign folder/package name."))

    @staticmethod
    def _merge_rules(rules: list[CampaignRule]) -> list[CampaignRule]:
        list_fields = {
            "allowed_countries", "allowed_industries", "job_title_keywords", "job_functions",
            "company_size", "revenue", "required_fields",
        }
        merged: list[CampaignRule] = []
        list_index: dict[str, CampaignRule] = {}
        for rule in rules:
            if rule.field not in list_fields:
                merged.append(rule)
                continue
            key = rule.key
            current = list_index.get(key)
            if current is None:
                current = CampaignRule(
                    rule.category, rule.field, rule.operator, list(rule.value or []),
                    confidence=rule.confidence, status=rule.status, notes=rule.notes,
                    sources=list(rule.sources),
                )
                list_index[key] = current
                merged.append(current)
                continue
            seen = {str(v).casefold() for v in current.value}
            for value in rule.value or []:
                if str(value).casefold() not in seen:
                    current.value.append(value)
                    seen.add(str(value).casefold())
            current.sources.extend(rule.sources)
            current.confidence = min(current.confidence, rule.confidence)
        return merged

    @staticmethod
    def _find_conflicts(rules: list[CampaignRule]) -> list[CampaignConflict]:
        groups: dict[str, list[CampaignRule]] = {}
        list_fields = {
            "allowed_countries", "allowed_industries", "job_title_keywords", "job_functions",
            "company_size", "revenue", "required_fields",
        }
        for rule in rules:
            if rule.field in list_fields:
                continue
            groups.setdefault(rule.key, []).append(rule)
        conflicts: list[CampaignConflict] = []
        for key, items in groups.items():
            distinct: dict[str, object] = {}
            sources: list[str] = []
            for rule in items:
                signature = json.dumps(rule.value, sort_keys=True, default=str).casefold()
                distinct[signature] = rule.value
                sources.extend(source.file for source in rule.sources)
            if len(distinct) > 1:
                conflicts.append(CampaignConflict(key, list(distinct.values()), sorted(set(sources))))
                for rule in items:
                    rule.status = "Needs Review"
                    if not rule.notes:
                        rule.notes = "Conflicting values were extracted from campaign sources."
        return conflicts


class CampaignReferenceLoader:
    """Load campaign-classified reference files with headerless-list support."""

    @staticmethod
    def read_rows(file_info: CampaignFileInfo) -> tuple[list[str], list[dict[str, str]]]:
        path = Path(file_info.absolute_path)
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm", ".xls"}:
            if suffix == ".xls":
                import xlrd
                book = xlrd.open_workbook(str(path), on_demand=True)
                sheet = book.sheet_by_index(0)
                if not sheet.nrows:
                    return [], []
                headers = [str(sheet.cell_value(0, col)).strip() for col in range(sheet.ncols)]
                rows = []
                for r in range(1, sheet.nrows):
                    rows.append({headers[c]: str(sheet.cell_value(r, c) or "") for c in range(sheet.ncols)})
                return headers, rows
            wb = load_workbook(path, data_only=True, read_only=True)
            ws = wb.active
            values = list(ws.iter_rows(values_only=True))
            if not values:
                return [], []
            headers = [str(v or "").strip() for v in values[0]]
            rows = [dict(zip(headers, ["" if v is None else str(v) for v in row])) for row in values[1:]]
            return headers, rows

        text = None
        encoding = "utf-8-sig"
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = path.read_text(encoding=enc)
                encoding = enc
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError(f"Could not decode {path.name}")
        lines = [line for line in text.splitlines() if line.strip()]
        if not lines:
            return [], []
        try:
            dialect = csv.Sniffer().sniff("\n".join(lines[:25]), delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
        parsed = list(csv.reader(lines, delimiter=delimiter))
        if file_info.metadata.get("headerless"):
            inferred = CampaignReferenceLoader._infer_header(file_info.role, parsed[0] if parsed else [])
            headers = inferred
            data_rows = parsed
        else:
            headers = [str(v).strip() for v in parsed[0]]
            data_rows = parsed[1:]
        rows = []
        for values in data_rows:
            padded = list(values) + [""] * max(0, len(headers) - len(values))
            rows.append({headers[i]: padded[i].strip() for i in range(len(headers))})
        return headers, rows

    @staticmethod
    def _infer_header(role: str, first_row: list[str]) -> list[str]:
        if len(first_row) <= 1:
            if role in {"client_suppression", "internal_suppression", "processed_suppression", "competitor_exclusion"}:
                return ["Email"]
            return ["Domain"]
        return [f"Column {i + 1}" for i in range(len(first_row))]
