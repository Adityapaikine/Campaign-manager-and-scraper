"""Import saved extraction results without changing evidence or existing edits."""
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from app.constants import STANDARD_HEADERS
from app.utils.text_utils import clean_missing


@dataclass(frozen=True)
class TransferResult:
    added: int
    skipped: int
    source: Path


def identity(row):
    value = clean_missing(row.get("Prospect LinkedIn Link"))
    try:
        parsed = urlsplit(value)
        host = (parsed.hostname or "").lower()
        if host == "linkedin.com" or host.endswith(".linkedin.com"):
            if parsed.path.startswith(("/in/", "/sales/lead/")):
                # Sales IDs are case sensitive; public profile slugs are not.
                path = parsed.path.rstrip("/")
                if path.startswith("/in/"):
                    path = path.casefold()
                return ("profile", path)
    except ValueError:
        pass
    # Without a reliable profile identity, only collapse identical content.
    ignored = {"Date", "Source File Name", "Campaign Name"}
    return ("record", *(clean_missing(row.get(h)) for h in STANDARD_HEADERS if h not in ignored))


def read_scraper_records(path):
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Choose a saved Scraper Excel workbook (.xlsx or .xlsm).")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Leads" not in book.sheetnames:
            raise ValueError("This workbook has no Leads sheet. Use Import in Leads Workspace for other files.")
        iterator = book["Leads"].iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(iterator, ())]
        if len(headers) != len(set(headers)) or not set(STANDARD_HEADERS).issubset(headers):
            raise ValueError("This workbook does not contain the complete Scraper lead columns.")
        rows = []
        for values in iterator:
            mapped = dict(zip(headers, values))
            row = {h: clean_missing(mapped.get(h)) for h in STANDARD_HEADERS}
            if all(v == "-" for v in row.values()):
                continue
            rows.append(row)
        return rows
    finally:
        book.close()


def transfer_workbook(controller, path):
    source = Path(path).resolve()
    rows = read_scraper_records(source)  # Validate the whole file before changing state.
    seen = {identity(r) for r in controller.data_manager.records}
    pending = []
    skipped = 0
    for row in rows:
        key = identity(row)
        if key in seen:
            skipped += 1
            continue
        seen.add(key)
        if controller.active_campaign and row["Campaign Name"] == "-":
            row["Campaign Name"] = controller.active_campaign.campaign_name
        pending.append(row)
    if pending:
        controller.data_manager.extend(pending, source_name=source.name)
        controller.clear_derived_state()
    return TransferResult(len(pending), skipped, source)
